#!/usr/bin/env python
"""Build a contextual, editable DOCX of core transport-planning assumptions.

Edit ``docs/advisor_consultation_questionnaire_content.json`` and rerun this
script. The output contains one title, two context paragraphs, and one
two-column Word table.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import tempfile
from pathlib import Path
from typing import Any, Iterable, Sequence
from zipfile import ZipFile

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONTENT = ROOT / "docs" / "advisor_consultation_questionnaire_content.json"
DEFAULT_OUTPUT = ROOT / "docs" / "전시_예비군_수송_연구_자문질문지.docx"

PRESET_NAME = "compact_reference_guide"
PAGE_WIDTH_IN = 8.5
PAGE_HEIGHT_IN = 11.0
MARGIN_IN = 1.0
HEADER_FOOTER_DISTANCE_IN = 0.492
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
TABLE_WIDTHS_DXA = (2400, 6960)
# Named override: preserve 10 pt table text while fitting title, context, and
# six decision rows on one page. Only vertical cell padding is tightened.
CELL_MARGINS_DXA = {"top": 60, "bottom": 60, "start": 120, "end": 120}

LATIN_FONT = "Calibri"
KOREAN_FONT = "Malgun Gothic"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "17212B"
HEADER_FILL = "E8EEF5"
BORDER = "C8D2DD"

COORDINATE_PAIR = re.compile(
    r"(?:3[3-9]\.[0-9]{4,})\s*[,/]\s*(?:12[4-9]\.[0-9]{4,})"
)
NONPUBLIC_DETAIL_TERMS = (
    "22사단",
    "22nd Infantry",
    "학야리",
    "실제 부대 좌표",
    "실시간 이동일정",
)
REMOVED_PACKAGING = (
    "자문 목적과 응답방법",
    "응답 기록 기준",
    "현재 모형과 자문 전 수정안",
    "1차 필수 자문항목",
    "2차 보충 자문항목",
    "토론용 기준·범위 초안",
    "최종 우선순위 질문",
    "회의 후 가정등록부",
    "자문 후 연구 반영 순서",
)
INTERNAL_ONLY_TERMS = (
    ".csv",
    ".json",
    ".yaml",
    ".py",
    "policy_id",
    "scenario_id",
    "mean_makespan",
    "CRN",
    "A1",
    "A2",
    "S→R",
    "600명",
    "30분 배차간격",
    "23대",
)


def _all_strings(value: Any) -> Iterable[str]:
    if isinstance(value, str):
        yield value
    elif isinstance(value, dict):
        for child in value.values():
            yield from _all_strings(child)
    elif isinstance(value, list):
        for child in value:
            yield from _all_strings(child)


def load_content(path: Path) -> dict[str, Any]:
    """Load and validate the intentionally small editable content schema."""

    with path.open("r", encoding="utf-8") as handle:
        content = json.load(handle)
    if not isinstance(content, dict):
        raise ValueError("questionnaire content must be a JSON object")
    expected_keys = {"schema_version", "editing_note", "title", "context", "items"}
    if set(content) != expected_keys:
        raise ValueError(f"content keys must be exactly {sorted(expected_keys)}")
    if content["schema_version"] != 3:
        raise ValueError("unsupported questionnaire schema_version")
    if not str(content["title"]).strip():
        raise ValueError("title must not be empty")
    context = content["context"]
    if not isinstance(context, list) or len(context) != 2:
        raise ValueError("context must contain exactly 2 paragraphs")
    if any(not isinstance(paragraph, str) or not paragraph.strip() for paragraph in context):
        raise ValueError("context paragraphs must be non-empty strings")

    items = content["items"]
    if not isinstance(items, list) or len(items) != 6:
        raise ValueError("items must contain exactly 6 core decisions")
    ids: list[str] = []
    titles: list[str] = []
    for index, item in enumerate(items, start=1):
        if not isinstance(item, dict) or set(item) != {"id", "title", "details"}:
            raise ValueError(f"item {index} must contain id, title, and details only")
        item_id = str(item["id"]).strip()
        title = str(item["title"]).strip()
        details = item["details"]
        if not item_id or not title:
            raise ValueError(f"item {index} id and title must not be empty")
        if not isinstance(details, list) or len(details) != 3:
            raise ValueError(f"item {index} must contain exactly 3 core details")
        if any(not isinstance(detail, str) or not detail.strip() for detail in details):
            raise ValueError(f"item {index} details must be non-empty strings")
        ids.append(item_id)
        titles.append(title)
    if len(ids) != len(set(ids)) or len(titles) != len(set(titles)):
        raise ValueError("item IDs and titles must be unique")

    rendered_content = {
        "title": content["title"],
        "context": content["context"],
        "items": content["items"],
    }
    joined = "\n".join(_all_strings(rendered_content))
    if COORDINATE_PAIR.search(joined):
        raise ValueError("content must not include precise coordinate pairs")
    for term in NONPUBLIC_DETAIL_TERMS:
        if term in joined:
            raise ValueError(f"content includes disallowed operational detail: {term}")
    for text in REMOVED_PACKAGING:
        if text in joined:
            raise ValueError(f"content includes removed document packaging: {text}")
    for text in INTERNAL_ONLY_TERMS:
        if text in joined:
            raise ValueError(f"content includes internal-only term: {text}")
    return content


def _hex_color(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#"))


def _set_font_refs(element, *, latin: str = LATIN_FONT, east_asia: str = KOREAN_FONT) -> None:
    r_pr = element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:cs"), east_asia)


def set_run_font(
    run,
    *,
    size: float,
    color: str = INK,
    bold: bool = False,
) -> None:
    run.font.name = LATIN_FONT
    run.font.size = Pt(size)
    run.font.color.rgb = _hex_color(color)
    run.bold = bold
    _set_font_refs(run._element)


def _style(document: DocxDocument, name: str):
    try:
        return document.styles[name]
    except KeyError:
        return document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def _configure_style(
    style,
    *,
    size: float,
    color: str,
    bold: bool,
    before: float,
    after: float,
    line_spacing: float,
) -> None:
    style.font.name = LATIN_FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = _hex_color(color)
    _set_font_refs(style.element)
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing
    fmt.widow_control = True


def configure_styles(document: DocxDocument) -> None:
    """Apply the compact_reference_guide token map."""

    _configure_style(
        document.styles["Normal"],
        size=11,
        color=INK,
        bold=False,
        before=0,
        after=6,
        line_spacing=1.25,
    )
    _configure_style(
        document.styles["Heading 1"],
        size=16,
        color=BLUE,
        bold=True,
        before=18,
        after=10,
        line_spacing=1.0,
    )
    _configure_style(
        _style(document, "Research Context"),
        size=10.2,
        color=INK,
        bold=False,
        before=0,
        after=6,
        line_spacing=1.2,
    )
    _configure_style(
        _style(document, "Core Item"),
        size=10.2,
        color=DARK_BLUE,
        bold=True,
        before=0,
        after=0,
        line_spacing=1.15,
    )
    _configure_style(
        _style(document, "Core Detail"),
        size=10.0,
        color=INK,
        bold=False,
        before=0,
        after=1,
        line_spacing=1.15,
    )


def configure_page(document: DocxDocument) -> None:
    section = document.sections[0]
    section.page_width = Inches(PAGE_WIDTH_IN)
    section.page_height = Inches(PAGE_HEIGHT_IN)
    section.top_margin = Inches(MARGIN_IN)
    section.right_margin = Inches(MARGIN_IN)
    section.bottom_margin = Inches(MARGIN_IN)
    section.left_margin = Inches(MARGIN_IN)
    section.header_distance = Inches(HEADER_FOOTER_DISTANCE_IN)
    section.footer_distance = Inches(HEADER_FOOTER_DISTANCE_IN)
    section.different_first_page_header_footer = False
    section.header.paragraphs[0].clear()
    section.footer.paragraphs[0].clear()


def _next_numbering_id(numbering, tag: str, attribute: str) -> int:
    values = [
        int(element.get(qn(attribute)))
        for element in numbering.findall(qn(tag))
        if element.get(qn(attribute)) is not None
    ]
    return max(values, default=-1) + 1


def create_bullet_numbering(document: DocxDocument) -> int:
    """Create valid compact-reference bullets before concrete num elements."""

    numbering = document.part.numbering_part.element
    abstract_id = _next_numbering_id(numbering, "w:abstractNum", "w:abstractNumId")
    num_id = _next_numbering_id(numbering, "w:num", "w:numId")

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"{0xA7100000 + abstract_id:08X}")
    abstract.append(nsid)
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    template = OxmlElement("w:tmpl")
    template.set(qn("w:val"), f"{0xB8200000 + abstract_id:08X}")
    abstract.append(template)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_format = OxmlElement("w:numFmt")
    num_format.set(qn("w:val"), "bullet")
    level.append(num_format)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "\uf0b7")
    level.append(level_text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    p_pr.append(indent)
    level.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Symbol")
    r_fonts.set(qn("w:hAnsi"), "Symbol")
    r_pr.append(r_fonts)
    level.append(r_pr)
    abstract.append(level)

    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def _apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_ref = OxmlElement("w:numId")
    num_ref.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_ref)
    p_pr.append(num_pr)


def _ensure_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def _set_width(parent, tag: str, width_dxa: int) -> None:
    element = _ensure_child(parent, tag)
    element.set(qn("w:type"), "dxa")
    element.set(qn("w:w"), str(int(width_dxa)))


def _set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = _ensure_child(tc_pr, "w:tcMar")
    for side, width in CELL_MARGINS_DXA.items():
        margin = _ensure_child(tc_mar, f"w:{side}")
        margin.set(qn("w:w"), str(width))
        margin.set(qn("w:type"), "dxa")


def _set_cell_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = _ensure_child(tc_pr, "w:tcBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = _ensure_child(borders, f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), BORDER)


def _shade_cell(cell, fill: str) -> None:
    shading = _ensure_child(cell._tc.get_or_add_tcPr(), "w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def _repeat_header(row) -> None:
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(marker)


def apply_table_geometry(table, widths_dxa: Sequence[int]) -> None:
    widths = [int(value) for value in widths_dxa]
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError("table widths must sum to CONTENT_WIDTH_DXA")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    _set_width(tbl_pr, "w:tblW", CONTENT_WIDTH_DXA)
    indent = _ensure_child(tbl_pr, "w:tblInd")
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    layout = _ensure_child(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for index, width in enumerate(widths):
        table.columns[index].width = Twips(width)
    for row in table.rows:
        row.height = None
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for index, cell in enumerate(row.cells):
            cell.width = Twips(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            _set_width(tc_pr, "w:tcW", widths[index])
            _set_cell_margins(cell)
            _set_cell_borders(cell)


def _set_header_cell(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.3, color=INK, bold=True)
    _shade_cell(cell, HEADER_FILL)


def _set_item_cell(cell, number: int, title: str) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.style = "Core Item"
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(f"{number}. {title}")
    set_run_font(run, size=10.2, color=DARK_BLUE, bold=True)


def _set_detail_cell(cell, details: Sequence[str], bullet_num_id: int) -> None:
    for index, detail in enumerate(details):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.clear()
        paragraph.style = "Core Detail"
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(1 if index < len(details) - 1 else 0)
        paragraph.paragraph_format.line_spacing = 1.15
        _apply_numbering(paragraph, bullet_num_id)
        run = paragraph.add_run(detail)
        set_run_font(run, size=10.0, color=INK)


def build_document(content: dict[str, Any]) -> DocxDocument:
    document = Document()
    configure_page(document)
    configure_styles(document)
    bullet_num_id = create_bullet_numbering(document)

    properties = document.core_properties
    properties.title = content["title"]
    properties.subject = "전시 동원예비군 수송 연구 핵심 계획가정"
    properties.author = "transport-system-sim research project"
    properties.comments = (
        "Generated by scripts/build_advisor_consultation_questionnaire.py; "
        f"preset={PRESET_NAME}; layout=title-context-table; "
        "override=context_table_compaction"
    )

    title = document.add_paragraph(style="Heading 1")
    title.paragraph_format.keep_with_next = True
    run = title.add_run(content["title"])
    set_run_font(run, size=16, color=BLUE, bold=True)

    for paragraph_text in content["context"]:
        paragraph = document.add_paragraph(style="Research Context")
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(paragraph_text)
        set_run_font(run, size=10.2, color=INK)

    table = document.add_table(rows=1 + len(content["items"]), cols=2)
    table.style = "Table Grid"
    _set_header_cell(table.rows[0].cells[0], "정해야 할 핵심 항목")
    _set_header_cell(table.rows[0].cells[1], "세부 결정내용")
    _repeat_header(table.rows[0])
    for number, item in enumerate(content["items"], start=1):
        _set_item_cell(table.rows[number].cells[0], number, item["title"])
        _set_detail_cell(table.rows[number].cells[1], item["details"], bullet_num_id)
    apply_table_geometry(table, TABLE_WIDTHS_DXA)
    return document


def _audit_table_geometry(path: Path) -> None:
    import xml.etree.ElementTree as ET

    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    attr = lambda name: f"{{{namespace['w']}}}{name}"
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    tables = root.findall(".//w:tbl", namespace)
    if len(tables) != 1:
        raise ValueError(f"expected one table, found {len(tables)}")
    table = tables[0]
    width_element = table.find("w:tblPr/w:tblW", namespace)
    indent_element = table.find("w:tblPr/w:tblInd", namespace)
    grid = [
        int(column.get(attr("w"), "0"))
        for column in table.findall("w:tblGrid/w:gridCol", namespace)
    ]
    if width_element is None or width_element.get(attr("type")) != "dxa":
        raise ValueError("table lacks DXA width")
    if int(width_element.get(attr("w"), "0")) != CONTENT_WIDTH_DXA:
        raise ValueError("table width is incorrect")
    if grid != list(TABLE_WIDTHS_DXA):
        raise ValueError("table grid is incorrect")
    if indent_element is None or int(indent_element.get(attr("w"), "-1")) != TABLE_INDENT_DXA:
        raise ValueError("table indent is incorrect")
    for row_index, row in enumerate(table.findall("w:tr", namespace), start=1):
        widths = []
        for cell in row.findall("w:tc", namespace):
            cell_width = cell.find("w:tcPr/w:tcW", namespace)
            widths.append(int(cell_width.get(attr("w"), "0")))
        if widths != list(TABLE_WIDTHS_DXA):
            raise ValueError(f"table row {row_index} widths do not match grid")


def _audit_reopened_document(path: Path, content: dict[str, Any]) -> None:
    document = Document(path)
    text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )
    expected = (
        [content["title"]]
        + list(content["context"])
        + [item["title"] for item in content["items"]]
    )
    missing = [value for value in expected if value not in text]
    if missing:
        raise ValueError(f"generated document is missing text: {missing}")
    leaked = [value for value in REMOVED_PACKAGING if value in text]
    if leaked:
        raise ValueError(f"generated document retains removed packaging: {leaked}")
    internal = [value for value in INTERNAL_ONLY_TERMS if value in text]
    if internal:
        raise ValueError(f"generated document retains internal-only terms: {internal}")
    if len(document.tables) != 1 or len(document.tables[0].rows) != 7:
        raise ValueError("generated document must contain one seven-row table")
    section = document.sections[0]
    expected_twips = {
        "page_width": 12240,
        "page_height": 15840,
        "top_margin": 1440,
        "right_margin": 1440,
        "bottom_margin": 1440,
        "left_margin": 1440,
    }
    for attribute, expected_value in expected_twips.items():
        actual = int(round(getattr(section, attribute).twips))
        if actual != expected_value:
            raise ValueError(f"{attribute}={actual}, expected {expected_value}")
    if section.header.paragraphs[0].text or section.footer.paragraphs[0].text:
        raise ValueError("minimal document must not have running header or footer text")
    if document.settings.element.find(qn("w:documentProtection")) is not None:
        raise ValueError("document must remain editable")
    _audit_table_geometry(path)


def save_document_atomic(
    document: DocxDocument,
    output_path: Path,
    content: dict[str, Any],
    *,
    force: bool,
) -> Path:
    output_path = output_path.resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists() and not force:
        raise FileExistsError(
            f"refusing to overwrite existing DOCX: {output_path}; pass --force after "
            "preserving edits in the JSON content"
        )
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}.", suffix=".docx", dir=output_path.parent
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        document.save(temporary_path)
        _audit_reopened_document(temporary_path, content)
        os.replace(temporary_path, output_path)
    finally:
        temporary_path.unlink(missing_ok=True)
    return output_path


def build_questionnaire(
    content_path: Path = DEFAULT_CONTENT,
    output_path: Path = DEFAULT_OUTPUT,
    *,
    force: bool = False,
) -> Path:
    content = load_content(content_path.resolve())
    document = build_document(content)
    return save_document_atomic(document, output_path, content, force=force)


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build the minimal Korean core-assumption DOCX."
    )
    parser.add_argument("--content", type=Path, default=DEFAULT_CONTENT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--force",
        action="store_true",
        help="overwrite an existing DOCX after preserving edits in JSON",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        output = build_questionnaire(args.content, args.output, force=args.force)
    except (FileNotFoundError, FileExistsError, ValueError, json.JSONDecodeError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    print(f"saved {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
