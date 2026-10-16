"""Guarded infrastructure for evidence-backed KIIE manuscript updates.

This module intentionally contains no manuscript prose.  It verifies the
publishable full-analysis bundle, exposes typed CSV tables, performs edits by
unique text anchors, and saves through a sibling temporary DOCX only after
protected-content invariants survive a reopen.
"""

from __future__ import annotations

import argparse
import csv
from copy import deepcopy
from dataclasses import dataclass
import hashlib
import json
import math
import os
from pathlib import Path
import re
from types import MappingProxyType
from typing import Any, Callable, Literal, Mapping, Sequence, TypeAlias
from uuid import uuid4

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


class ManuscriptUpdateError(RuntimeError):
    """Base class for deterministic manuscript-update failures."""


class AnalysisIntegrityError(ManuscriptUpdateError):
    """Raised when analysis evidence is incomplete, stale, or modified."""


class AnchorNotFoundError(ManuscriptUpdateError):
    """Raised when a requested paragraph anchor does not exist."""


class AnchorNotUniqueError(ManuscriptUpdateError):
    """Raised when a requested paragraph anchor is ambiguous."""


class ProtectedContentError(ManuscriptUpdateError):
    """Raised when author, drawing, media, or section invariants change."""


CsvScalar: TypeAlias = str | int | float | bool | None | dict[str, Any] | list[Any]
AnchorMode: TypeAlias = Literal["exact", "contains"]


_INT_COLUMNS = frozenset(
    {
        "arrival_seed",
        "threat_seed",
        "threat_draw",
        "corridor_path_count",
        "demand",
        "road_fleet_total",
        "morris_point",
        "morris_trajectory",
        "selected_edge_count",
        "analysis_graph_nodes",
        "analysis_graph_edges",
        "censored_count",
        "vehicle_cycles",
        "road_vehicle_cycles",
        "road_deployed_seat_capacity",
        "road_boarded_passengers",
        "rail_deployed_seat_capacity",
        "rail_boarded_passengers",
        "empty_return_trips",
        "train_trips",
        "graph_node_count",
        "graph_edge_count",
        "path_edge_count",
        "path_diversity_limit",
        "path_diversity_count",
        "full_path_edge_count",
        "total_pair_count",
        "completion_pair_count",
        "finite_pair_count",
        "excluded_pair_count",
        "bracket_count",
        "joint_seed_bootstrap_complete_seed_count",
        "joint_seed_bootstrap_replicates",
        "joint_seed_bootstrap_successful_replicates",
        "joint_seed_bootstrap_seed",
        "normal_time_match_count",
        "connectivity_match_count",
        "policy_ranking_match_count",
        "bootstrap_replicates",
        "bootstrap_seed",
        "bootstrap_n",
        "t_n",
        "completion_t_n",
        "completion_bootstrap_n",
        "completion_bootstrap_replicates",
        "completion_bootstrap_seed",
        "replicate_count",
        "seed_max",
        "threat_draw_count",
        "finite_observation_count",
        "bootstrap_excluded_observation_count",
        "completion_delta_bootstrap_replicates",
        "completion_delta_bootstrap_seed",
        "completion_threat_draw_count",
        "completion_observation_count",
        "probability_proxy_observation_count",
        "trajectory_count",
        "sample_point_count",
        "expected_arrival_seed_count",
        "incomplete_group_count",
        "missing_seed_count",
        "nonfinite_seed_count",
        "run_count",
        "valid_completion_count",
        "full_completion_count",
        "finite_makespan_count",
        "rank",
    }
)

_FLOAT_COLUMNS = frozenset(
    {
        "completion_rate",
        "makespan",
        "makespan_min",
        "penalized_makespan",
        "rail_multiplier",
        "road_multiplier",
        "road_travel_time_multiplier",
        "disruption_capacity_factor",
        "road_longhaul_multiplier",
        "rail_travel_multiplier",
        "transfer_time_min",
        "arrival_sigma",
        "empty_return_minutes",
        "mean_vehicle_load_factor",
        "road_mean_vehicle_load_factor",
        "rail_mean_load_factor",
        "assembly_wait_minutes",
        "transfer_wait_minutes",
        "rail_wait_minutes",
        "shortest_travel_time_min",
        "shortest_distance_m",
        "full_path_edge_overlap_ratio",
        "full_path_edge_jaccard",
        "travel_time_detour_ratio_vs_full",
        "distance_detour_ratio_vs_full",
        "mean_bus_completion_rate",
        "mean_multimodal_completion_rate",
        "mean_bus_makespan",
        "mean_multimodal_makespan",
        "mean_completion_rate_delta_bus_minus_multimodal",
        "completion_t_ci_lower",
        "completion_t_ci_upper",
        "completion_t_confidence",
        "completion_t_standard_error",
        "completion_bootstrap_ci_lower",
        "completion_bootstrap_ci_upper",
        "completion_bootstrap_confidence",
        "mean_delta_bus_minus_multimodal",
        "t_ci_lower",
        "t_ci_upper",
        "t_confidence",
        "t_standard_error",
        "bootstrap_ci_lower",
        "bootstrap_ci_upper",
        "bootstrap_confidence",
        "crossing_road_multiplier",
        "lower_road_multiplier",
        "upper_road_multiplier",
        "lower_mean_delta",
        "upper_mean_delta",
        "joint_seed_bootstrap_ci_lower",
        "joint_seed_bootstrap_ci_upper",
        "joint_seed_bootstrap_confidence",
        "normal_time_mean_absolute_error",
        "normal_time_mean_absolute_percentage_error",
        "connectivity_agreement",
        "policy_ranking_agreement",
        "completion_delta_bootstrap_ci_lower",
        "completion_delta_bootstrap_ci_upper",
        "completion_delta_bootstrap_confidence",
        "bus_positive_completion_probability_proxy",
        "bus_positive_completion_proxy_ci_lower",
        "bus_positive_completion_proxy_ci_upper",
        "multimodal_positive_completion_probability_proxy",
        "multimodal_positive_completion_proxy_ci_lower",
        "multimodal_positive_completion_proxy_ci_upper",
        "bus_full_completion_probability_proxy",
        "bus_full_completion_proxy_ci_lower",
        "bus_full_completion_proxy_ci_upper",
        "multimodal_full_completion_probability_proxy",
        "multimodal_full_completion_proxy_ci_lower",
        "multimodal_full_completion_proxy_ci_upper",
        "probability_proxy_confidence",
        "mu",
        "mu_star",
        "sigma",
        "mu_star_conf",
        "t_ci_half_width",
        "estimate_change_from_previous",
        "half_width_ratio_to_previous",
        "mean_completion_rate",
        "mean_finite_makespan",
    }
)

_BOOL_COLUMNS = frozenset({"connected", "full_reference_connected", "winner"})
_JSON_COLUMNS = frozenset(
    {
        "completion_incomplete_reason_counts",
        "incomplete_reason_counts",
        "joint_seed_bootstrap_excluded_seeds",
        "empty_threat_draws",
        "completion_empty_threat_draws",
        "completion_incomplete_reason_counts",
    }
)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _is_sha256(value: object) -> bool:
    return (
        isinstance(value, str)
        and len(value) == 64
        and all(character in "0123456789abcdefABCDEF" for character in value)
    )


def _csv_data_row_count(path: Path) -> int:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        try:
            next(reader)
        except StopIteration as error:
            raise AnalysisIntegrityError(f"CSV has no header: {path}") from error
        return sum(1 for _ in reader)


def _resolve_relative(root: Path, raw: object, label: str) -> Path:
    if not isinstance(raw, str) or not raw.strip():
        raise AnalysisIntegrityError(f"{label} path must be non-empty text")
    relative = Path(raw)
    if relative.is_absolute():
        raise AnalysisIntegrityError(f"{label} path must be relative: {raw}")
    resolved_root = root.resolve()
    resolved = (resolved_root / relative).resolve()
    try:
        resolved.relative_to(resolved_root)
    except ValueError as error:
        raise AnalysisIntegrityError(f"{label} path escapes evidence root: {raw}") from error
    return resolved


def _verify_csv_record(root: Path, raw: object, label: str) -> Path:
    if not isinstance(raw, Mapping):
        raise AnalysisIntegrityError(f"{label} artifact must be an object")
    path = _resolve_relative(root, raw.get("path"), label)
    if path.suffix.lower() != ".csv":
        raise AnalysisIntegrityError(f"{label} artifact must be CSV: {path}")
    if not path.is_file():
        raise AnalysisIntegrityError(f"{label} CSV is missing: {path}")
    recorded_hash = raw.get("sha256")
    if not _is_sha256(recorded_hash) or str(recorded_hash).lower() != _sha256(path):
        raise AnalysisIntegrityError(f"{label} CSV SHA-256 mismatch: {path}")
    recorded_rows = raw.get("data_row_count")
    if isinstance(recorded_rows, bool) or not isinstance(recorded_rows, int):
        raise AnalysisIntegrityError(f"{label} data_row_count must be an integer: {path}")
    actual_rows = _csv_data_row_count(path)
    if recorded_rows != actual_rows:
        raise AnalysisIntegrityError(
            f"{label} CSV row count mismatch: {path} ({recorded_rows} != {actual_rows})"
        )
    return path


def _parse_typed_cell(column: str, raw: str, path: Path, row_number: int) -> CsvScalar:
    value = raw.strip()
    if value == "":
        return None
    try:
        if column in _INT_COLUMNS:
            if not re.fullmatch(r"[+-]?\d+", value):
                raise ValueError("not an integer")
            return int(value)
        if column in _FLOAT_COLUMNS:
            number = float(value)
            if not math.isfinite(number):
                raise ValueError("not finite")
            return number
        if column in _BOOL_COLUMNS:
            normalized = value.lower()
            if normalized not in {"true", "false"}:
                raise ValueError("not a boolean")
            return normalized == "true"
        if column in _JSON_COLUMNS:
            parsed = json.loads(value)
            if not isinstance(parsed, (dict, list)):
                raise ValueError("not a JSON object or array")
            return parsed
    except (ValueError, json.JSONDecodeError) as error:
        raise AnalysisIntegrityError(
            f"invalid {column!r} value in {path} row {row_number}: {value!r}"
        ) from error
    return value


@dataclass(frozen=True)
class TypedCsvTable:
    """Hash-verified analysis CSV with schema-aware scalar conversion."""

    name: str
    path: Path
    fieldnames: tuple[str, ...]
    rows: tuple[Mapping[str, CsvScalar], ...]
    sha256: str

    def select(self, **criteria: object) -> tuple[Mapping[str, CsvScalar], ...]:
        unknown = set(criteria) - set(self.fieldnames)
        if unknown:
            raise KeyError(f"unknown {self.name} columns: {sorted(unknown)}")
        return tuple(
            row
            for row in self.rows
            if all(row[column] == expected for column, expected in criteria.items())
        )

    def one(self, **criteria: object) -> Mapping[str, CsvScalar]:
        matches = self.select(**criteria)
        if len(matches) != 1:
            raise AnalysisIntegrityError(
                f"{self.name} expected one row for {criteria}, found {len(matches)}"
            )
        return matches[0]


def _read_typed_csv(path: Path, *, name: str, sha256: str) -> TypedCsvTable:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None or any(not field for field in reader.fieldnames):
            raise AnalysisIntegrityError(f"CSV has invalid header: {path}")
        if len(set(reader.fieldnames)) != len(reader.fieldnames):
            raise AnalysisIntegrityError(f"CSV has duplicate columns: {path}")
        rows: list[Mapping[str, CsvScalar]] = []
        for row_number, raw in enumerate(reader, start=2):
            parsed = {
                column: _parse_typed_cell(
                    column, "" if raw_value is None else raw_value, path, row_number
                )
                for column, raw_value in raw.items()
                if column is not None
            }
            rows.append(MappingProxyType(parsed))
    return TypedCsvTable(
        name=name,
        path=path,
        fieldnames=tuple(reader.fieldnames),
        rows=tuple(rows),
        sha256=sha256,
    )


@dataclass(frozen=True)
class VerifiedAnalysis:
    """Validated full-analysis manifest and its typed generated tables."""

    manifest_path: Path
    output_root: Path
    analysis_directory: Path
    manifest: Mapping[str, Any]
    tables: Mapping[str, TypedCsvTable]

    def table(self, name: str) -> TypedCsvTable:
        key = Path(name).stem
        try:
            return self.tables[key]
        except KeyError as error:
            raise KeyError(
                f"analysis table {name!r} not found; available={sorted(self.tables)}"
            ) from error


def load_verified_analysis(manifest_path: str | Path) -> VerifiedAnalysis:
    """Load only a complete, blocker-free full analysis and verify every CSV."""

    path = Path(manifest_path).expanduser().resolve()
    try:
        raw_manifest = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as error:
        raise AnalysisIntegrityError(f"cannot read analysis manifest {path}: {error}") from error
    if not isinstance(raw_manifest, dict):
        raise AnalysisIntegrityError("analysis manifest must contain an object")
    if raw_manifest.get("stage") != "full":
        raise AnalysisIntegrityError("analysis manifest must set stage=full")
    if raw_manifest.get("analysis_complete_full") is not True:
        raise AnalysisIntegrityError(
            "analysis manifest must set analysis_complete_full=true"
        )
    blockers = raw_manifest.get("analysis_completeness_blockers")
    if not isinstance(blockers, list) or blockers:
        raise AnalysisIntegrityError(
            "analysis manifest analysis_completeness_blockers must be empty"
        )
    if raw_manifest.get("final_study_ready") is not False:
        raise AnalysisIntegrityError("analysis manifest must preserve final_study_ready=false")

    expected_analysis_dir = path.parent
    expected_output_root = path.parents[2]
    try:
        recorded_analysis_dir = Path(raw_manifest["analysis_directory"]).expanduser().resolve()
        recorded_output_root = Path(raw_manifest["output_root"]).expanduser().resolve()
    except (KeyError, TypeError) as error:
        raise AnalysisIntegrityError("analysis manifest lacks evidence directories") from error
    if recorded_analysis_dir != expected_analysis_dir:
        raise AnalysisIntegrityError("analysis_directory does not match manifest location")
    if recorded_output_root != expected_output_root:
        raise AnalysisIntegrityError("output_root does not match manifest location")

    analyses = raw_manifest.get("analyses")
    if not isinstance(analyses, Mapping) or not analyses:
        raise AnalysisIntegrityError("analysis manifest lacks analysis status records")
    unavailable = [
        name
        for name, state in analyses.items()
        if not isinstance(state, Mapping) or state.get("status") != "available"
    ]
    if unavailable:
        raise AnalysisIntegrityError(f"analysis status is not available: {sorted(unavailable)}")

    source_files = raw_manifest.get("source_files")
    if not isinstance(source_files, list) or not source_files:
        raise AnalysisIntegrityError("analysis manifest lacks source CSV records")
    for index, record in enumerate(source_files):
        _verify_csv_record(expected_output_root, record, f"source_files[{index}]")

    generated = raw_manifest.get("generated_csv_files")
    if not isinstance(generated, list) or not generated:
        raise AnalysisIntegrityError("analysis manifest lacks generated CSV records")
    tables: dict[str, TypedCsvTable] = {}
    seen_paths: set[Path] = set()
    for index, record in enumerate(generated):
        generated_path = _verify_csv_record(
            expected_analysis_dir, record, f"generated_csv_files[{index}]"
        )
        if generated_path in seen_paths:
            raise AnalysisIntegrityError(f"duplicate generated CSV record: {generated_path}")
        seen_paths.add(generated_path)
        table_name = generated_path.stem
        if table_name in tables:
            raise AnalysisIntegrityError(f"duplicate generated table name: {table_name}")
        tables[table_name] = _read_typed_csv(
            generated_path,
            name=table_name,
            sha256=str(record["sha256"]).lower(),
        )

    return VerifiedAnalysis(
        manifest_path=path,
        output_root=expected_output_root,
        analysis_directory=expected_analysis_dir,
        manifest=MappingProxyType(raw_manifest),
        tables=MappingProxyType(tables),
    )


def _paragraph_matches(paragraph: Paragraph, anchor: str, mode: AnchorMode) -> bool:
    text = paragraph.text.strip()
    normalized_anchor = anchor.strip()
    if mode == "exact":
        return text == normalized_anchor
    if mode == "contains":
        return normalized_anchor in text
    raise ValueError("anchor mode must be exact or contains")


def find_unique_paragraph(
    document: DocxDocument,
    anchor: str,
    *,
    mode: AnchorMode = "exact",
    style_name: str | None = None,
) -> Paragraph:
    """Return exactly one paragraph selected by human-readable text anchor."""

    if not isinstance(anchor, str) or not anchor.strip():
        raise ValueError("paragraph anchor must be non-empty text")
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if _paragraph_matches(paragraph, anchor, mode)
        and (style_name is None or paragraph.style.name == style_name)
    ]
    if not matches:
        raise AnchorNotFoundError(f"paragraph anchor not found: {anchor!r}")
    if len(matches) != 1:
        raise AnchorNotUniqueError(
            f"paragraph anchor is not unique: {anchor!r} ({len(matches)} matches)"
        )
    return matches[0]


def _next_table_element(paragraph: Paragraph) -> object | None:
    element = paragraph._p.getnext()
    while element is not None:
        if element.tag.endswith("}tbl"):
            return element
        if element.tag.endswith("}sectPr"):
            return None
        if element.tag.endswith("}p"):
            visible_text = "".join(
                node.text or "" for node in element.xpath(".//w:t")
            ).strip()
            if visible_text:
                return None
        element = element.getnext()
    return None


def find_table_after_anchor(
    document: DocxDocument,
    anchor: str,
    *,
    mode: AnchorMode = "exact",
    style_name: str | None = None,
) -> Table:
    """Find first body table following one unique caption/heading anchor."""

    paragraph = find_unique_paragraph(
        document, anchor, mode=mode, style_name=style_name
    )
    table_element = _next_table_element(paragraph)
    if table_element is None:
        raise AnchorNotFoundError(f"no table follows paragraph anchor: {anchor!r}")
    for table in document.tables:
        if table._tbl is table_element:
            return table
    raise AnchorNotFoundError(f"following table is outside document body: {anchor!r}")


def _base_run_properties(paragraph: Paragraph) -> object | None:
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return deepcopy(run._r.rPr)
    return None


def _paragraph_contains_drawing(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing | .//w:pict"))


def replace_paragraph_text(
    paragraph: Paragraph,
    text: object,
    *,
    first_run_label: str | None = None,
    first_run_bold: bool = True,
) -> Paragraph:
    """Replace text while leaving paragraph properties/style untouched."""

    replacement = str(text)
    if _paragraph_contains_drawing(paragraph):
        raise ProtectedContentError("refusing to replace paragraph containing a drawing")
    if first_run_label is not None and not replacement.startswith(first_run_label):
        raise ValueError("replacement text must begin with first_run_label")
    base_rpr = _base_run_properties(paragraph)
    for child in list(paragraph._p):
        if child is not paragraph._p.pPr:
            paragraph._p.remove(child)

    fragments = (
        [(first_run_label, first_run_bold), (replacement[len(first_run_label) :], None)]
        if first_run_label is not None
        else [(replacement, None)]
    )
    for fragment, bold in fragments:
        if not fragment:
            continue
        run = paragraph.add_run(fragment)
        if base_rpr is not None:
            existing = run._r.rPr
            if existing is not None:
                run._r.remove(existing)
            run._r.insert(0, deepcopy(base_rpr))
        if bold is not None:
            run.bold = bold
    return paragraph


def _cell_value_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("table cell float must be finite")
        return format(value, ".15g")
    return str(value)


def _replace_cell_text(cell: _Cell, value: object) -> None:
    paragraphs = cell.paragraphs
    paragraph = paragraphs[0]
    replace_paragraph_text(paragraph, _cell_value_text(value))
    for extra in paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def _strip_repeating_header(row_element: object) -> None:
    for header in row_element.xpath("./w:trPr/w:tblHeader"):
        header.getparent().remove(header)


_NUMERIC_SECTION_HEADING_PATTERN = re.compile(
    r"^\d+(?:\.\d+)*\.?\s+\S"
)


def _set_keep_with_next(paragraph_element: object) -> None:
    paragraph_properties = paragraph_element.get_or_add_pPr()
    for existing in paragraph_properties.xpath("./w:keepNext"):
        paragraph_properties.remove(existing)
    paragraph_properties.get_or_add_keepNext()


def apply_pagination_controls(document: DocxDocument) -> DocxDocument:
    """Apply deterministic anti-orphan controls to manuscript body content."""

    for table_element in document.element.body.xpath(".//w:tbl"):
        rows = table_element.xpath("./w:tr")
        for row_index, row_element in enumerate(rows):
            row_properties = row_element.get_or_add_trPr()
            for tag in ("cantSplit", "tblHeader"):
                for existing in row_properties.xpath(f"./w:{tag}"):
                    row_properties.remove(existing)
            row_properties.append(OxmlElement("w:cantSplit"))
            if row_index == 0:
                row_properties.append(OxmlElement("w:tblHeader"))
        if rows:
            for paragraph_element in rows[0].xpath("./w:tc/w:p"):
                _set_keep_with_next(paragraph_element)

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("<표 ") or _NUMERIC_SECTION_HEADING_PATTERN.match(text):
            _set_keep_with_next(paragraph._p)
    return document


def replace_table_rows(
    table: Table,
    rows: Sequence[Sequence[object]],
    *,
    preserve_leading_rows: int = 0,
) -> Table:
    """Replace table rows without replacing tblPr, tblGrid, or cell tcPr."""

    if isinstance(preserve_leading_rows, bool) or preserve_leading_rows < 0:
        raise ValueError("preserve_leading_rows must be a non-negative integer")
    current_rows = len(table.rows)
    if preserve_leading_rows > current_rows:
        raise ValueError("cannot preserve more rows than table contains")
    column_count = len(table.columns)
    normalized_rows = [tuple(row) for row in rows]
    for row_number, row in enumerate(normalized_rows, start=1):
        if len(row) != column_count:
            raise ValueError(
                f"replacement row {row_number} has {len(row)} cells; expected {column_count}"
            )
    target_count = preserve_leading_rows + len(normalized_rows)
    if target_count == 0:
        raise ValueError("a Word table must retain at least one row")

    while len(table.rows) < target_count:
        template_index = (
            preserve_leading_rows
            if len(table.rows) > preserve_leading_rows
            else max(0, preserve_leading_rows - 1)
        )
        template = deepcopy(table.rows[template_index]._tr)
        _strip_repeating_header(template)
        table._tbl.append(template)
    while len(table.rows) > target_count:
        table._tbl.remove(table.rows[-1]._tr)

    for row_offset, values in enumerate(normalized_rows):
        target_row = table.rows[preserve_leading_rows + row_offset]
        for cell, value in zip(target_row.cells, values):
            _replace_cell_text(cell, value)
    return table


def _canonical_xml_hash(element: object) -> str:
    from lxml import etree

    encoded = etree.tostring(element, method="c14n", exclusive=True)
    return hashlib.sha256(encoded).hexdigest()


def _section_signature(document: DocxDocument) -> tuple[tuple[object, ...], ...]:
    fields = (
        "start_type",
        "orientation",
        "page_width",
        "page_height",
        "top_margin",
        "bottom_margin",
        "left_margin",
        "right_margin",
        "header_distance",
        "footer_distance",
        "gutter",
    )
    return tuple(
        tuple(getattr(section, field) for field in fields)
        for section in document.sections
    )


def _drawing_signatures(document: DocxDocument) -> tuple[tuple[str, tuple[str, ...]], ...]:
    signatures: list[tuple[str, tuple[str, ...]]] = []
    for part in document.part.package.parts:
        element = getattr(part, "element", None)
        if element is None:
            continue
        drawings = element.xpath(".//w:drawing | .//w:pict")
        if drawings:
            signatures.append(
                (
                    str(part.partname),
                    tuple(_canonical_xml_hash(drawing) for drawing in drawings),
                )
            )
    return tuple(sorted(signatures))


def _image_part_signatures(document: DocxDocument) -> tuple[tuple[str, str], ...]:
    return tuple(
        sorted(
            (
                str(part.partname),
                hashlib.sha256(part.blob).hexdigest(),
            )
            for part in document.part.package.parts
            if str(getattr(part, "content_type", "")).startswith("image/")
        )
    )


@dataclass(frozen=True)
class DocumentInvariants:
    protected_paragraph_indices: tuple[int, ...]
    protected_paragraph_texts: tuple[str, ...]
    protected_paragraph_hashes: tuple[str, ...]
    drawing_signatures: tuple[tuple[str, tuple[str, ...]], ...]
    image_part_signatures: tuple[tuple[str, str], ...]
    inline_shape_count: int
    section_signature: tuple[tuple[object, ...], ...]
    styles_signature: str

    @classmethod
    def capture(
        cls,
        document: DocxDocument,
        *,
        protected_paragraph_indices: Sequence[int] = (3, 4, 5),
    ) -> "DocumentInvariants":
        indices = tuple(protected_paragraph_indices)
        if any(isinstance(index, bool) or index < 0 for index in indices):
            raise ValueError("protected paragraph indices must be non-negative integers")
        if any(index >= len(document.paragraphs) for index in indices):
            raise ProtectedContentError("protected author paragraph is missing")
        paragraphs = [document.paragraphs[index] for index in indices]
        return cls(
            protected_paragraph_indices=indices,
            protected_paragraph_texts=tuple(paragraph.text for paragraph in paragraphs),
            protected_paragraph_hashes=tuple(
                _canonical_xml_hash(paragraph._p) for paragraph in paragraphs
            ),
            drawing_signatures=_drawing_signatures(document),
            image_part_signatures=_image_part_signatures(document),
            inline_shape_count=len(document.inline_shapes),
            section_signature=_section_signature(document),
            styles_signature=_canonical_xml_hash(document.styles.element),
        )

    def assert_preserved(self, document: DocxDocument) -> None:
        if any(index >= len(document.paragraphs) for index in self.protected_paragraph_indices):
            raise ProtectedContentError("protected author paragraph is missing")
        paragraphs = [
            document.paragraphs[index] for index in self.protected_paragraph_indices
        ]
        texts = tuple(paragraph.text for paragraph in paragraphs)
        hashes = tuple(_canonical_xml_hash(paragraph._p) for paragraph in paragraphs)
        if texts != self.protected_paragraph_texts or hashes != self.protected_paragraph_hashes:
            raise ProtectedContentError("protected author paragraph changed")
        if _drawing_signatures(document) != self.drawing_signatures:
            raise ProtectedContentError("unrelated drawing XML changed")
        if _image_part_signatures(document) != self.image_part_signatures:
            raise ProtectedContentError("unrelated drawing media changed")
        if len(document.inline_shapes) != self.inline_shape_count:
            raise ProtectedContentError("unrelated inline drawing count changed")
        if _section_signature(document) != self.section_signature:
            raise ProtectedContentError("document section geometry changed")
        if _canonical_xml_hash(document.styles.element) != self.styles_signature:
            raise ProtectedContentError("document style definitions changed")


@dataclass
class ManuscriptSession:
    source_path: Path
    source_sha256: str
    document: DocxDocument
    invariants: DocumentInvariants


def open_manuscript_session(
    source_path: str | Path,
    *,
    protected_paragraph_indices: Sequence[int] = (3, 4, 5),
) -> ManuscriptSession:
    """Open one manuscript and freeze non-negotiable content invariants."""

    path = Path(source_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(path)
    document = Document(path)
    return ManuscriptSession(
        source_path=path,
        source_sha256=_sha256(path),
        document=document,
        invariants=DocumentInvariants.capture(
            document, protected_paragraph_indices=protected_paragraph_indices
        ),
    )


def atomic_save_document(
    session: ManuscriptSession,
    output_path: str | Path,
    *,
    validate: Callable[[DocxDocument], None] | None = None,
) -> DocxDocument:
    """Save, reopen, validate, then atomically replace destination DOCX."""

    target = Path(output_path).expanduser().resolve()
    if target.suffix.lower() != ".docx":
        raise ValueError("output path must end in .docx")
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.parent / f".{target.stem}.{uuid4().hex}.tmp.docx"
    try:
        session.invariants.assert_preserved(session.document)
        session.document.save(temporary)
        reopened_temporary = Document(temporary)
        session.invariants.assert_preserved(reopened_temporary)
        if validate is not None:
            validate(reopened_temporary)
        os.replace(temporary, target)
        reopened_target = Document(target)
        session.invariants.assert_preserved(reopened_target)
        if validate is not None:
            validate(reopened_target)
        return reopened_target
    finally:
        if temporary.exists():
            temporary.unlink()


def _following_table_exists(paragraph: Paragraph) -> bool:
    return _next_table_element(paragraph) is not None


def build_document_inventory(
    document: DocxDocument,
    *,
    anchors: Sequence[str] = (),
) -> dict[str, Any]:
    """Return a read-only structural inventory for dry-run planning."""

    nonempty = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    counts: dict[str, int] = {}
    for text in nonempty:
        counts[text] = counts.get(text, 0) + 1
    anchor_inventory: dict[str, dict[str, Any]] = {}
    for anchor in anchors:
        matches = [
            paragraph
            for paragraph in document.paragraphs
            if _paragraph_matches(paragraph, anchor, "exact")
        ]
        anchor_inventory[anchor] = {
            "match_count": len(matches),
            "style_names": [paragraph.style.name for paragraph in matches],
            "following_table": (
                len(matches) == 1 and _following_table_exists(matches[0])
            ),
        }
    return {
        "paragraph_count": len(document.paragraphs),
        "nonempty_paragraph_count": len(nonempty),
        "table_count": len(document.tables),
        "tables": [
            {
                "row_count": len(table.rows),
                "column_count": len(table.columns),
                "style_name": table.style.name if table.style is not None else None,
            }
            for table in document.tables
        ],
        "inline_shape_count": len(document.inline_shapes),
        "section_count": len(document.sections),
        "duplicate_exact_paragraphs": {
            text: count for text, count in sorted(counts.items()) if count > 1
        },
        "anchors": anchor_inventory,
    }


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Inventory manuscript structure and verify full revision evidence."
    )
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--analysis-manifest", type=Path)
    parser.add_argument("--anchor", action="append", default=[])
    parser.add_argument("--dry-run", action="store_true")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    if not args.dry_run:
        raise SystemExit("no content plan supplied; use --dry-run or import guarded APIs")
    session = open_manuscript_session(args.input)
    inventory = build_document_inventory(session.document, anchors=args.anchor)
    inventory["input_path"] = str(session.source_path)
    inventory["input_sha256"] = session.source_sha256
    if args.analysis_manifest is not None:
        analysis = load_verified_analysis(args.analysis_manifest)
        inventory["analysis"] = {
            "manifest_path": str(analysis.manifest_path),
            "stage": analysis.manifest["stage"],
            "analysis_complete_full": analysis.manifest["analysis_complete_full"],
            "analysis_completeness_blockers": list(
                analysis.manifest["analysis_completeness_blockers"]
            ),
            "tables": {
                name: {
                    "path": str(table.path),
                    "row_count": len(table.rows),
                    "column_count": len(table.fieldnames),
                    "sha256": table.sha256,
                }
                for name, table in sorted(analysis.tables.items())
            },
        }
    print(json.dumps(inventory, ensure_ascii=False, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
