"""Declarative, evidence-bound revision plans for the KIIE manuscript.

Plans are JSON documents pinned to one source DOCX and one verified full-analysis
manifest by SHA-256.  Every target is resolved before mutation.  Numeric values
can be materialized from hash-verified analysis CSVs; surrounding prose remains
an explicit author-owned literal in the plan.  Saving delegates to the guarded
atomic DOCX writer in :mod:`update_revision_manuscript`.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
import hashlib
import json
import math
from pathlib import Path
import re
import sys
from types import MappingProxyType
from typing import Any, Literal, Mapping, Sequence, TypeAlias

from docx.document import Document as DocxDocument

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from paper._build.update_revision_manuscript import (
    AnalysisIntegrityError,
    AnchorMode,
    CsvScalar,
    ManuscriptSession,
    VerifiedAnalysis,
    apply_pagination_controls,
    atomic_save_document,
    find_table_after_anchor,
    find_unique_paragraph,
    load_verified_analysis,
    open_manuscript_session,
    replace_paragraph_text,
    replace_table_rows,
)


class ContentPlanError(RuntimeError):
    """Raised when a content plan is invalid, stale, or cannot resolve evidence."""


JsonScalar: TypeAlias = str | int | float | bool | None
OperationKind: TypeAlias = Literal["paragraph", "table_rows"]

_PLAN_ID_PATTERN = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.-]{0,127}$")
_SHA256_PATTERN = re.compile(r"^[0-9a-fA-F]{64}$")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _object(value: object, label: str) -> Mapping[str, Any]:
    if not isinstance(value, Mapping):
        raise ContentPlanError(f"{label} must be an object")
    return value


def _keys(
    value: Mapping[str, Any],
    label: str,
    *,
    required: Sequence[str],
    optional: Sequence[str] = (),
) -> None:
    missing = set(required) - set(value)
    unknown = set(value) - set(required) - set(optional)
    if missing:
        raise ContentPlanError(f"{label} missing keys: {sorted(missing)}")
    if unknown:
        raise ContentPlanError(f"{label} has unsupported keys: {sorted(unknown)}")


def _text(value: object, label: str, *, allow_empty: bool = False) -> str:
    if not isinstance(value, str) or (not allow_empty and not value.strip()):
        qualifier = "text" if allow_empty else "non-empty text"
        raise ContentPlanError(f"{label} must be {qualifier}")
    return value


def _hash_text(value: object, label: str) -> str:
    text = _text(value, label)
    if not _SHA256_PATTERN.fullmatch(text):
        raise ContentPlanError(f"{label} must be a 64-character SHA-256")
    return text.lower()


def _json_scalar(value: object, label: str) -> JsonScalar:
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float) and math.isfinite(value):
        return value
    raise ContentPlanError(f"{label} must be a finite JSON scalar")


def _where(value: object, label: str) -> Mapping[str, JsonScalar]:
    raw = _object(value, label)
    converted: dict[str, JsonScalar] = {}
    for key, item in raw.items():
        column = _text(key, f"{label} column")
        converted[column] = _json_scalar(item, f"{label}.{column}")
    return MappingProxyType(converted)


def _format_spec(value: object, label: str) -> str:
    spec = _text(value, label, allow_empty=True)
    if len(spec) > 64 or "{" in spec or "}" in spec:
        raise ContentPlanError(f"{label} is not a safe format specification")
    return spec


@dataclass(frozen=True)
class AnchorSpec:
    text: str
    mode: AnchorMode = "exact"
    style_name: str | None = None


@dataclass(frozen=True)
class EvidenceReference:
    table: str
    where: Mapping[str, JsonScalar]
    column: str
    format_spec: str = ""
    prefix: str = ""
    suffix: str = ""


@dataclass(frozen=True)
class LiteralFragment:
    literal: str


@dataclass(frozen=True)
class EvidenceFragment:
    evidence: EvidenceReference


ParagraphFragment: TypeAlias = LiteralFragment | EvidenceFragment


@dataclass(frozen=True)
class ParagraphOperation:
    operation_id: str
    kind: Literal["paragraph"]
    anchor: AnchorSpec
    fragments: tuple[ParagraphFragment, ...]
    first_run_label: str | None = None
    first_run_bold: bool = True


@dataclass(frozen=True)
class TableSource:
    table: str
    where: Mapping[str, JsonScalar]
    sort_by: tuple[str, ...]


@dataclass(frozen=True)
class LiteralCell:
    literal: str


@dataclass(frozen=True)
class ColumnCell:
    column: str
    format_spec: str = ""
    prefix: str = ""
    suffix: str = ""


TableCell: TypeAlias = LiteralCell | ColumnCell


@dataclass(frozen=True)
class TableRowsOperation:
    operation_id: str
    kind: Literal["table_rows"]
    anchor: AnchorSpec
    preserve_leading_rows: int
    source: TableSource | None
    cells: tuple[TableCell, ...]
    literal_rows: tuple[tuple[str, ...], ...] = ()
    header: tuple[str, ...] | None = None


PlanOperation: TypeAlias = ParagraphOperation | TableRowsOperation


@dataclass(frozen=True)
class ContentPlan:
    path: Path
    sha256: str
    schema_version: int
    plan_id: str
    source_sha256: str
    analysis_manifest_sha256: str
    operations: tuple[PlanOperation, ...]
    description: str | None = None


@dataclass(frozen=True)
class _PreparedParagraph:
    operation: ParagraphOperation
    paragraph_index: int
    resolved_text: str


@dataclass(frozen=True)
class _PreparedTable:
    operation: TableRowsOperation
    table_index: int
    resolved_rows: tuple[tuple[str, ...], ...]
    expected_rows_after_apply: tuple[tuple[str, ...], ...]


PreparedOperation: TypeAlias = _PreparedParagraph | _PreparedTable


@dataclass(frozen=True)
class _PreparedPlan:
    session: ManuscriptSession
    analysis: VerifiedAnalysis
    plan: ContentPlan
    operations: tuple[PreparedOperation, ...]


def _parse_anchor(value: object, label: str) -> AnchorSpec:
    raw = _object(value, label)
    _keys(raw, label, required=("text",), optional=("mode", "style_name"))
    mode = raw.get("mode", "exact")
    if mode not in {"exact", "contains"}:
        raise ContentPlanError(f"{label}.mode must be exact or contains")
    style_name_raw = raw.get("style_name")
    style_name = (
        None
        if style_name_raw is None
        else _text(style_name_raw, f"{label}.style_name")
    )
    return AnchorSpec(
        text=_text(raw["text"], f"{label}.text"),
        mode=mode,
        style_name=style_name,
    )


def _parse_evidence(value: object, label: str) -> EvidenceReference:
    raw = _object(value, label)
    _keys(
        raw,
        label,
        required=("table", "where", "column"),
        optional=("format", "prefix", "suffix"),
    )
    return EvidenceReference(
        table=_text(raw["table"], f"{label}.table"),
        where=_where(raw["where"], f"{label}.where"),
        column=_text(raw["column"], f"{label}.column"),
        format_spec=_format_spec(raw.get("format", ""), f"{label}.format"),
        prefix=_text(raw.get("prefix", ""), f"{label}.prefix", allow_empty=True),
        suffix=_text(raw.get("suffix", ""), f"{label}.suffix", allow_empty=True),
    )


def _parse_fragment(value: object, label: str) -> ParagraphFragment:
    raw = _object(value, label)
    if set(raw) == {"literal"}:
        return LiteralFragment(_text(raw["literal"], f"{label}.literal", allow_empty=True))
    if set(raw) == {"evidence"}:
        return EvidenceFragment(_parse_evidence(raw["evidence"], f"{label}.evidence"))
    raise ContentPlanError(f"{label} must contain exactly literal or evidence")


def _parse_table_source(value: object, label: str) -> TableSource:
    raw = _object(value, label)
    _keys(raw, label, required=("table",), optional=("where", "sort_by"))
    sort_raw = raw.get("sort_by", [])
    if not isinstance(sort_raw, list):
        raise ContentPlanError(f"{label}.sort_by must be an array")
    sort_by = tuple(
        _text(item, f"{label}.sort_by[{index}]")
        for index, item in enumerate(sort_raw)
    )
    if len(set(sort_by)) != len(sort_by):
        raise ContentPlanError(f"{label}.sort_by contains duplicate columns")
    return TableSource(
        table=_text(raw["table"], f"{label}.table"),
        where=_where(raw.get("where", {}), f"{label}.where"),
        sort_by=sort_by,
    )


def _parse_cell(value: object, label: str) -> TableCell:
    raw = _object(value, label)
    if set(raw) == {"literal"}:
        return LiteralCell(_text(raw["literal"], f"{label}.literal", allow_empty=True))
    _keys(
        raw,
        label,
        required=("column",),
        optional=("format", "prefix", "suffix"),
    )
    return ColumnCell(
        column=_text(raw["column"], f"{label}.column"),
        format_spec=_format_spec(raw.get("format", ""), f"{label}.format"),
        prefix=_text(raw.get("prefix", ""), f"{label}.prefix", allow_empty=True),
        suffix=_text(raw.get("suffix", ""), f"{label}.suffix", allow_empty=True),
    )


def _parse_operation(value: object, index: int) -> PlanOperation:
    label = f"operations[{index}]"
    raw = _object(value, label)
    kind = raw.get("kind")
    operation_id = _text(raw.get("id"), f"{label}.id")
    if not _PLAN_ID_PATTERN.fullmatch(operation_id):
        raise ContentPlanError(f"{label}.id has unsupported characters")
    if kind == "paragraph":
        _keys(
            raw,
            label,
            required=("id", "kind", "anchor", "replacement"),
            optional=("first_run_label", "first_run_bold"),
        )
        replacement = _object(raw["replacement"], f"{label}.replacement")
        _keys(replacement, f"{label}.replacement", required=("fragments",))
        fragments_raw = replacement["fragments"]
        if not isinstance(fragments_raw, list) or not fragments_raw:
            raise ContentPlanError(f"{label}.replacement.fragments must be a non-empty array")
        first_label_raw = raw.get("first_run_label")
        first_label = (
            None
            if first_label_raw is None
            else _text(first_label_raw, f"{label}.first_run_label")
        )
        first_bold = raw.get("first_run_bold", True)
        if not isinstance(first_bold, bool):
            raise ContentPlanError(f"{label}.first_run_bold must be boolean")
        return ParagraphOperation(
            operation_id=operation_id,
            kind="paragraph",
            anchor=_parse_anchor(raw["anchor"], f"{label}.anchor"),
            fragments=tuple(
                _parse_fragment(fragment, f"{label}.replacement.fragments[{fragment_index}]")
                for fragment_index, fragment in enumerate(fragments_raw)
            ),
            first_run_label=first_label,
            first_run_bold=first_bold,
        )
    if kind == "table_rows":
        _keys(
            raw,
            label,
            required=("id", "kind", "anchor", "preserve_leading_rows"),
            optional=("source", "cells", "rows", "header"),
        )
        preserve = raw["preserve_leading_rows"]
        if isinstance(preserve, bool) or not isinstance(preserve, int) or preserve < 0:
            raise ContentPlanError(f"{label}.preserve_leading_rows must be non-negative integer")
        has_source = "source" in raw or "cells" in raw
        has_rows = "rows" in raw
        if has_source == has_rows:
            raise ContentPlanError(
                f"{label} must define exactly one of source+cells or rows"
            )
        if has_source and not ({"source", "cells"} <= set(raw)):
            raise ContentPlanError(f"{label} must define source and cells together")
        cells_raw = raw.get("cells", [])
        if has_source and (not isinstance(cells_raw, list) or not cells_raw):
            raise ContentPlanError(f"{label}.cells must be a non-empty array")
        literal_rows: tuple[tuple[str, ...], ...] = ()
        if has_rows:
            rows_raw = raw["rows"]
            if not isinstance(rows_raw, list) or not rows_raw:
                raise ContentPlanError(f"{label}.rows must be a non-empty array")
            parsed_rows: list[tuple[str, ...]] = []
            for row_index, row_raw in enumerate(rows_raw):
                if not isinstance(row_raw, list) or not row_raw:
                    raise ContentPlanError(
                        f"{label}.rows[{row_index}] must be a non-empty array"
                    )
                parsed_rows.append(
                    tuple(
                        _text(
                            value,
                            f"{label}.rows[{row_index}][{cell_index}]",
                            allow_empty=True,
                        )
                        for cell_index, value in enumerate(row_raw)
                    )
                )
            literal_rows = tuple(parsed_rows)
        header_raw = raw.get("header")
        if header_raw is not None:
            if not isinstance(header_raw, list) or not header_raw:
                raise ContentPlanError(f"{label}.header must be a non-empty array")
            header = tuple(
                _text(value, f"{label}.header[{cell_index}]", allow_empty=True)
                for cell_index, value in enumerate(header_raw)
            )
        else:
            header = None
        return TableRowsOperation(
            operation_id=operation_id,
            kind="table_rows",
            anchor=_parse_anchor(raw["anchor"], f"{label}.anchor"),
            preserve_leading_rows=preserve,
            source=(
                _parse_table_source(raw["source"], f"{label}.source")
                if has_source
                else None
            ),
            cells=tuple(
                _parse_cell(cell, f"{label}.cells[{cell_index}]")
                for cell_index, cell in enumerate(cells_raw)
            ),
            literal_rows=literal_rows,
            header=header,
        )
    raise ContentPlanError(f"{label}.kind must be paragraph or table_rows")


def load_content_plan(plan_path: str | Path) -> ContentPlan:
    """Parse one strict content plan without opening the manuscript or evidence."""

    path = Path(plan_path).expanduser().resolve()
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as error:
        raise ContentPlanError(f"cannot read content plan {path}: {error}") from error
    root = _object(raw, "content plan")
    _keys(
        root,
        "content plan",
        required=(
            "schema_version",
            "plan_id",
            "source_sha256",
            "analysis_manifest_sha256",
            "operations",
        ),
        optional=("description",),
    )
    if root["schema_version"] != 1 or isinstance(root["schema_version"], bool):
        raise ContentPlanError("content plan schema_version must equal 1")
    plan_id = _text(root["plan_id"], "content plan.plan_id")
    if not _PLAN_ID_PATTERN.fullmatch(plan_id):
        raise ContentPlanError("content plan.plan_id has unsupported characters")
    operations_raw = root["operations"]
    if not isinstance(operations_raw, list) or not operations_raw:
        raise ContentPlanError("content plan.operations must be a non-empty array")
    operations = tuple(
        _parse_operation(operation, index)
        for index, operation in enumerate(operations_raw)
    )
    ids = [operation.operation_id for operation in operations]
    if len(set(ids)) != len(ids):
        raise ContentPlanError("content plan contains duplicate operation ids")
    description_raw = root.get("description")
    description = (
        None
        if description_raw is None
        else _text(description_raw, "content plan.description", allow_empty=True)
    )
    return ContentPlan(
        path=path,
        sha256=_sha256(path),
        schema_version=1,
        plan_id=plan_id,
        source_sha256=_hash_text(root["source_sha256"], "content plan.source_sha256"),
        analysis_manifest_sha256=_hash_text(
            root["analysis_manifest_sha256"],
            "content plan.analysis_manifest_sha256",
        ),
        operations=operations,
        description=description,
    )


def _scalar_text(value: CsvScalar, format_spec: str, label: str) -> str:
    if value is None:
        rendered = ""
    elif format_spec:
        try:
            rendered = format(value, format_spec)
        except (TypeError, ValueError) as error:
            raise ContentPlanError(
                f"{label} cannot format {type(value).__name__} with {format_spec!r}"
            ) from error
    elif isinstance(value, (dict, list)):
        rendered = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    elif isinstance(value, float):
        if not math.isfinite(value):
            raise ContentPlanError(f"{label} is non-finite")
        rendered = format(value, ".15g")
    else:
        rendered = str(value)
    return rendered


def _analysis_table(analysis: VerifiedAnalysis, name: str, label: str):
    try:
        return analysis.table(name)
    except KeyError as error:
        raise ContentPlanError(f"{label} references unknown analysis table {name!r}") from error


def _validate_columns(table, columns: Sequence[str], label: str) -> None:
    missing = sorted(set(columns) - set(table.fieldnames))
    if missing:
        raise ContentPlanError(f"{label} references unknown columns: {missing}")


def _select_rows(table, where: Mapping[str, JsonScalar], label: str):
    _validate_columns(table, tuple(where), label)
    try:
        return table.select(**dict(where))
    except KeyError as error:
        raise ContentPlanError(f"{label} has invalid selection: {error}") from error


def _resolve_evidence(
    reference: EvidenceReference,
    analysis: VerifiedAnalysis,
    label: str,
) -> str:
    table = _analysis_table(analysis, reference.table, label)
    _validate_columns(table, (reference.column,), label)
    rows = _select_rows(table, reference.where, label)
    if len(rows) != 1:
        raise ContentPlanError(
            f"{label} expected exactly one evidence row, found {len(rows)}"
        )
    value = _scalar_text(rows[0][reference.column], reference.format_spec, label)
    return f"{reference.prefix}{value}{reference.suffix}"


def _sort_key(value: CsvScalar) -> tuple[int, str, object]:
    if value is None:
        return (1, "none", "")
    if isinstance(value, bool):
        return (0, "bool", int(value))
    if isinstance(value, (int, float)):
        return (0, "number", float(value))
    if isinstance(value, (dict, list)):
        rendered = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    else:
        rendered = str(value)
    return (0, type(value).__name__, rendered)


def _resolve_paragraph(
    operation: ParagraphOperation,
    analysis: VerifiedAnalysis,
) -> str:
    pieces: list[str] = []
    for index, fragment in enumerate(operation.fragments):
        if isinstance(fragment, LiteralFragment):
            pieces.append(fragment.literal)
        else:
            pieces.append(
                _resolve_evidence(
                    fragment.evidence,
                    analysis,
                    f"operation {operation.operation_id!r} fragment {index}",
                )
            )
    resolved = "".join(pieces)
    if not resolved:
        raise ContentPlanError(f"operation {operation.operation_id!r} resolves to empty text")
    if operation.first_run_label is not None and not resolved.startswith(
        operation.first_run_label
    ):
        raise ContentPlanError(
            f"operation {operation.operation_id!r} text does not begin with first_run_label"
        )
    return resolved


def _resolve_table_rows(
    operation: TableRowsOperation,
    analysis: VerifiedAnalysis,
) -> tuple[tuple[str, ...], ...]:
    label = f"operation {operation.operation_id!r}"
    if operation.literal_rows:
        return operation.literal_rows
    if operation.source is None:
        raise ContentPlanError(f"{label} has neither source rows nor literal rows")
    table = _analysis_table(analysis, operation.source.table, label)
    cell_columns = tuple(
        cell.column for cell in operation.cells if isinstance(cell, ColumnCell)
    )
    _validate_columns(table, (*operation.source.sort_by, *cell_columns), label)
    rows = list(_select_rows(table, operation.source.where, label))
    if not rows:
        raise ContentPlanError(f"{label} source selection returned no rows")
    for column in reversed(operation.source.sort_by):
        rows.sort(key=lambda row, name=column: _sort_key(row[name]))
    resolved: list[tuple[str, ...]] = []
    for row_index, row in enumerate(rows):
        cells: list[str] = []
        for cell_index, cell in enumerate(operation.cells):
            if isinstance(cell, LiteralCell):
                cells.append(cell.literal)
                continue
            value = _scalar_text(
                row[cell.column],
                cell.format_spec,
                f"{label} row {row_index} cell {cell_index}",
            )
            cells.append(f"{cell.prefix}{value}{cell.suffix}")
        resolved.append(tuple(cells))
    return tuple(resolved)


def _paragraph_index(document: DocxDocument, paragraph) -> int:
    for index, candidate in enumerate(document.paragraphs):
        if candidate._p is paragraph._p:
            return index
    raise ContentPlanError("resolved paragraph is outside document body")


def _table_index(document: DocxDocument, table) -> int:
    for index, candidate in enumerate(document.tables):
        if candidate._tbl is table._tbl:
            return index
    raise ContentPlanError("resolved table is outside document body")


def _prepare(
    source_path: str | Path,
    analysis_manifest_path: str | Path,
    plan_path: str | Path,
) -> _PreparedPlan:
    session = open_manuscript_session(source_path)
    analysis = load_verified_analysis(analysis_manifest_path)
    plan = load_content_plan(plan_path)
    if session.source_sha256 != plan.source_sha256:
        raise ContentPlanError(
            "content plan source_sha256 does not match input manuscript"
        )
    manifest_hash = _sha256(analysis.manifest_path)
    if manifest_hash != plan.analysis_manifest_sha256:
        raise ContentPlanError(
            "content plan analysis_manifest_sha256 does not match verified analysis"
        )

    prepared: list[PreparedOperation] = []
    target_keys: set[tuple[str, int]] = set()
    for operation in plan.operations:
        anchor = operation.anchor
        if isinstance(operation, ParagraphOperation):
            paragraph = find_unique_paragraph(
                session.document,
                anchor.text,
                mode=anchor.mode,
                style_name=anchor.style_name,
            )
            index = _paragraph_index(session.document, paragraph)
            target_key = ("paragraph", index)
            resolved = _resolve_paragraph(operation, analysis)
            prepared_operation: PreparedOperation = _PreparedParagraph(
                operation=operation,
                paragraph_index=index,
                resolved_text=resolved,
            )
        else:
            table = find_table_after_anchor(
                session.document,
                anchor.text,
                mode=anchor.mode,
                style_name=anchor.style_name,
            )
            index = _table_index(session.document, table)
            target_key = ("table", index)
            if operation.preserve_leading_rows > len(table.rows):
                raise ContentPlanError(
                    f"operation {operation.operation_id!r} preserves more rows than target contains"
                )
            if operation.source is not None and len(operation.cells) != len(table.columns):
                raise ContentPlanError(
                    f"operation {operation.operation_id!r} defines {len(operation.cells)} cells; "
                    f"target table has {len(table.columns)} columns"
                )
            if operation.header is not None:
                if operation.preserve_leading_rows < 1:
                    raise ContentPlanError(
                        f"operation {operation.operation_id!r} defines a header but preserves no rows"
                    )
                if len(operation.header) != len(table.columns):
                    raise ContentPlanError(
                        f"operation {operation.operation_id!r} defines {len(operation.header)} "
                        f"header cells; target table has {len(table.columns)} columns"
                    )
            resolved_rows = _resolve_table_rows(operation, analysis)
            wrong_width = [
                row_index
                for row_index, row in enumerate(resolved_rows, start=1)
                if len(row) != len(table.columns)
            ]
            if wrong_width:
                raise ContentPlanError(
                    f"operation {operation.operation_id!r} replacement rows have wrong "
                    f"column count at rows {wrong_width}"
                )
            leading_rows = [
                tuple(cell.text for cell in row.cells)
                for row in table.rows[: operation.preserve_leading_rows]
            ]
            if operation.header is not None:
                leading_rows[0] = operation.header
            leading = tuple(leading_rows)
            prepared_operation = _PreparedTable(
                operation=operation,
                table_index=index,
                resolved_rows=resolved_rows,
                expected_rows_after_apply=leading + resolved_rows,
            )
        if target_key in target_keys:
            raise ContentPlanError(
                f"content plan has duplicate target at operation {operation.operation_id!r}"
            )
        target_keys.add(target_key)
        prepared.append(prepared_operation)
    return _PreparedPlan(
        session=session,
        analysis=analysis,
        plan=plan,
        operations=tuple(prepared),
    )


def _rows_hash(rows: Sequence[Sequence[str]]) -> str:
    encoded = json.dumps(
        [list(row) for row in rows],
        ensure_ascii=False,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _preview(prepared: _PreparedPlan, *, preview_row_limit: int) -> dict[str, Any]:
    if isinstance(preview_row_limit, bool) or preview_row_limit < 0:
        raise ValueError("preview_row_limit must be a non-negative integer")
    operations: list[dict[str, Any]] = []
    for item in prepared.operations:
        operation = item.operation
        base: dict[str, Any] = {
            "id": operation.operation_id,
            "kind": operation.kind,
            "anchor": operation.anchor.text,
            "anchor_mode": operation.anchor.mode,
            "anchor_style_name": operation.anchor.style_name,
        }
        if isinstance(item, _PreparedParagraph):
            base.update(
                {
                    "paragraph_index": item.paragraph_index,
                    "resolved_text": item.resolved_text,
                }
            )
        else:
            preview_rows = item.resolved_rows[:preview_row_limit]
            base.update(
                {
                    "table_index": item.table_index,
                    "preserve_leading_rows": item.operation.preserve_leading_rows,
                    "resolved_row_count": len(item.resolved_rows),
                    "resolved_rows_sha256": _rows_hash(item.resolved_rows),
                    "resolved_rows_preview": [list(row) for row in preview_rows],
                    "resolved_rows_preview_truncated": len(preview_rows)
                    < len(item.resolved_rows),
                    "header": (
                        None
                        if item.operation.header is None
                        else list(item.operation.header)
                    ),
                }
            )
        operations.append(base)
    return {
        "dry_run": True,
        "schema_version": prepared.plan.schema_version,
        "plan_id": prepared.plan.plan_id,
        "plan_path": str(prepared.plan.path),
        "plan_sha256": prepared.plan.sha256,
        "source_path": str(prepared.session.source_path),
        "source_sha256": prepared.session.source_sha256,
        "analysis_manifest_path": str(prepared.analysis.manifest_path),
        "analysis_manifest_sha256": prepared.plan.analysis_manifest_sha256,
        "analysis_complete_full": prepared.analysis.manifest[
            "analysis_complete_full"
        ],
        "final_study_ready": prepared.analysis.manifest["final_study_ready"],
        "operation_count": len(prepared.operations),
        "operations": operations,
    }


def preview_content_plan(
    source_path: str | Path,
    analysis_manifest_path: str | Path,
    plan_path: str | Path,
    *,
    preview_row_limit: int = 5,
) -> dict[str, Any]:
    """Fully preflight a plan and return materialized edits without writing."""

    prepared = _prepare(source_path, analysis_manifest_path, plan_path)
    return _preview(prepared, preview_row_limit=preview_row_limit)


def _validate_applied(document: DocxDocument, operations: Sequence[PreparedOperation]) -> None:
    for item in operations:
        if isinstance(item, _PreparedParagraph):
            try:
                actual = document.paragraphs[item.paragraph_index].text
            except IndexError as error:
                raise ContentPlanError(
                    f"applied paragraph {item.operation.operation_id!r} is missing after reopen"
                ) from error
            if actual != item.resolved_text:
                raise ContentPlanError(
                    f"applied paragraph {item.operation.operation_id!r} changed after reopen"
                )
            continue
        try:
            table = document.tables[item.table_index]
        except IndexError as error:
            raise ContentPlanError(
                f"applied table {item.operation.operation_id!r} is missing after reopen"
            ) from error
        actual_rows = tuple(
            tuple(cell.text for cell in row.cells) for row in table.rows
        )
        if actual_rows != item.expected_rows_after_apply:
            raise ContentPlanError(
                f"applied table {item.operation.operation_id!r} changed after reopen"
            )


def apply_content_plan(
    source_path: str | Path,
    analysis_manifest_path: str | Path,
    plan_path: str | Path,
    output_path: str | Path,
    *,
    preview_row_limit: int = 5,
) -> dict[str, Any]:
    """Preflight every operation, apply to a private session, and save atomically."""

    prepared = _prepare(source_path, analysis_manifest_path, plan_path)
    report = _preview(prepared, preview_row_limit=preview_row_limit)

    for item in prepared.operations:
        if isinstance(item, _PreparedParagraph):
            paragraph = prepared.session.document.paragraphs[item.paragraph_index]
            replace_paragraph_text(
                paragraph,
                item.resolved_text,
                first_run_label=item.operation.first_run_label,
                first_run_bold=item.operation.first_run_bold,
            )
        else:
            table = prepared.session.document.tables[item.table_index]
            replace_table_rows(
                table,
                item.resolved_rows,
                preserve_leading_rows=item.operation.preserve_leading_rows,
            )
            if item.operation.header is not None:
                for cell, value in zip(table.rows[0].cells, item.operation.header):
                    replace_paragraph_text(cell.paragraphs[0], value)

    target = Path(output_path).expanduser().resolve()
    apply_pagination_controls(prepared.session.document)
    atomic_save_document(
        prepared.session,
        target,
        validate=lambda document: _validate_applied(document, prepared.operations),
    )
    report["dry_run"] = False
    report["output_path"] = str(target)
    report["output_sha256"] = _sha256(target)
    return report


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Preview or atomically apply an evidence-bound manuscript content plan."
    )
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--analysis-manifest", required=True, type=Path)
    parser.add_argument("--plan", required=True, type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--preview-row-limit", type=int, default=5)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    try:
        if args.dry_run:
            if args.output is not None:
                raise ContentPlanError("--output is not allowed with --dry-run")
            report = preview_content_plan(
                args.input,
                args.analysis_manifest,
                args.plan,
                preview_row_limit=args.preview_row_limit,
            )
        else:
            if args.output is None:
                raise ContentPlanError("--output is required when applying a plan")
            report = apply_content_plan(
                args.input,
                args.analysis_manifest,
                args.plan,
                args.output,
                preview_row_limit=args.preview_row_limit,
            )
    except (AnalysisIntegrityError, ContentPlanError) as error:
        raise SystemExit(str(error)) from error
    print(json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
