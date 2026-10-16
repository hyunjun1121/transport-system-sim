# -*- coding: utf-8 -*-
"""Embed figures next to known anchors using captions.md as caption source.

Only an image paragraph immediately preceding a known ``<그림 N>`` caption is
reused. Other drawings are untouched. Output is saved to a sibling temporary
DOCX, reopened and validated, then atomically moved into place.
"""

from __future__ import annotations

import argparse
import hashlib
import os
import re
import sys
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Mapping, Sequence

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Inches


sys.stdout.reconfigure(encoding="utf-8")

DOCX = Path("paper/KIIE_투고논문.docx")
CAPTIONS_PATH = Path("paper/figures/captions.md")


@dataclass(frozen=True)
class FigureSpec:
    number: int
    path: Path
    width_in: float

    @property
    def prefix(self) -> str:
        return f"<그림 {self.number}>"


FIGURES = (
    FigureSpec(1, Path("paper/figures/fig1_framework.png"), 6.3),
    FigureSpec(2, Path("paper/figures/fig2_corridor_map.png"), 5.0),
    FigureSpec(3, Path("paper/figures/fig3_bpr_noop.png"), 5.4),
    FigureSpec(4, Path("paper/figures/fig4_segment_decomposition.png"), 6.7),
    FigureSpec(5, Path("paper/figures/fig5_rail_substitution.png"), 5.4),
    FigureSpec(6, Path("paper/figures/fig6_assumption_collapse.png"), 6.1),
)

_CAPTION_START = re.compile(r"^\s*\*\*그림\s+(\d+)\.\*\*\s*(.*)\s*$")


def load_captions(path: Path, required_numbers: Iterable[int]) -> dict[str, str]:
    """Parse Markdown figure captions and convert their labels to DOCX anchors."""
    path = Path(path)
    lines = path.read_text(encoding="utf-8").splitlines()
    parsed: dict[int, list[str]] = {}
    current: int | None = None

    for line in lines:
        match = _CAPTION_START.match(line)
        if match:
            number = int(match.group(1))
            if number in parsed:
                raise ValueError(f"duplicate caption: 그림 {number}")
            current = number
            parsed[number] = [match.group(2).strip()]
            continue
        if current is None:
            continue
        stripped = line.strip()
        if stripped == "---" or stripped.startswith("#"):
            current = None
        elif stripped:
            parsed[current].append(stripped)

    required = tuple(required_numbers)
    missing = [number for number in required if number not in parsed]
    if missing:
        raise ValueError("missing captions: " + ", ".join(f"그림 {number}" for number in missing))

    captions = {}
    for number in required:
        body = " ".join(part for part in parsed[number] if part).strip()
        if not body:
            raise ValueError(f"empty caption: 그림 {number}")
        prefix = f"<그림 {number}>"
        captions[prefix] = f"{prefix} {body}"
    return captions


def _has_drawing(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.findall(".//" + qn("w:drawing")))


def _is_pure_single_picture_paragraph(paragraph: Paragraph) -> bool:
    """Return true only when clearing the paragraph cannot remove visible text."""
    drawings = paragraph._p.findall(".//" + qn("w:drawing"))
    blips = paragraph._p.xpath(".//a:blip")
    return paragraph.text.strip() == "" and len(drawings) == 1 and len(blips) == 1


def _previous_paragraph(paragraph: Paragraph) -> Paragraph | None:
    previous = paragraph._p.getprevious()
    if previous is None or previous.tag != qn("w:p"):
        return None
    return Paragraph(previous, paragraph._parent)


def _clear_content_preserve_paragraph_properties(paragraph: Paragraph) -> None:
    properties = paragraph._p.pPr
    for child in list(paragraph._p):
        if child is not properties:
            paragraph._p.remove(child)


def _replace_text_preserve_paragraph_format(paragraph: Paragraph, text: str) -> None:
    """Replace visible text while retaining pPr, style and first-run formatting."""
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _find_caption(document: DocumentType, prefix: str) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.strip().startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"expected one {prefix} anchor, found {len(matches)}")
    return matches[0]


def _target_image_paragraph(caption: Paragraph) -> Paragraph:
    previous = _previous_paragraph(caption)
    if previous is not None and (
        _is_pure_single_picture_paragraph(previous)
        or previous.text.strip() == "[ 그림 삽입 예정 ]"
    ):
        _clear_content_preserve_paragraph_properties(previous)
        return previous

    image_paragraph = caption.insert_paragraph_before("")
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.keep_with_next = True
    return image_paragraph


def _section_signature(document: DocumentType) -> tuple[tuple[object, ...], ...]:
    fields = (
        "orientation",
        "page_width",
        "page_height",
        "top_margin",
        "right_margin",
        "bottom_margin",
        "left_margin",
        "header_distance",
        "footer_distance",
        "gutter",
        "start_type",
    )
    return tuple(tuple(getattr(section, field) for field in fields) for section in document.sections)


def _validate_saved_document(
    path: Path,
    captions: Mapping[str, str],
    figures: Sequence[FigureSpec],
    section_signature: tuple[tuple[object, ...], ...],
) -> None:
    reopened = Document(path)
    if _section_signature(reopened) != section_signature:
        raise RuntimeError("section formatting changed during figure embedding")
    for spec in figures:
        caption = _find_caption(reopened, spec.prefix)
        if caption.text != captions[spec.prefix]:
            raise RuntimeError(f"caption mismatch after save: {spec.prefix}")
        image_paragraph = _previous_paragraph(caption)
        if image_paragraph is None or not _has_drawing(image_paragraph):
            raise RuntimeError(f"missing adjacent drawing after save: {spec.prefix}")
        blips = image_paragraph._p.xpath(".//a:blip")
        if len(blips) != 1:
            raise RuntimeError(f"expected one adjacent image after save: {spec.prefix}")
        relationship_id = blips[0].get(qn("r:embed"))
        if not relationship_id or relationship_id not in image_paragraph.part.related_parts:
            raise RuntimeError(f"missing image relationship after save: {spec.prefix}")
        embedded_blob = image_paragraph.part.related_parts[relationship_id].blob
        source_blob = Path(spec.path).read_bytes()
        if len(embedded_blob) != len(source_blob) or hashlib.sha256(embedded_blob).digest() != hashlib.sha256(
            source_blob
        ).digest():
            raise RuntimeError(f"image blob mismatch after save: {spec.prefix}")


def _validate_figure_specs(figures: Sequence[FigureSpec]) -> None:
    seen: set[int] = set()
    for spec in figures:
        if spec.number in seen:
            raise ValueError(f"duplicate figure number: {spec.number}")
        seen.add(spec.number)


def embed_figures(
    *,
    docx_path: Path = DOCX,
    captions_path: Path = CAPTIONS_PATH,
    output_path: Path | None = None,
    figures: Sequence[FigureSpec] = FIGURES,
) -> Path:
    """Update known figure pairs and atomically write a validated DOCX."""
    source = Path(docx_path)
    destination = Path(output_path) if output_path is not None else source
    figures = tuple(figures)
    _validate_figure_specs(figures)
    captions = load_captions(Path(captions_path), (spec.number for spec in figures))
    missing_images = [str(spec.path) for spec in figures if not Path(spec.path).is_file()]
    if missing_images:
        raise FileNotFoundError("missing figure images: " + ", ".join(missing_images))

    document = Document(source)
    original_sections = _section_signature(document)
    for spec in figures:
        caption = _find_caption(document, spec.prefix)
        _replace_text_preserve_paragraph_format(caption, captions[spec.prefix])
        image_paragraph = _target_image_paragraph(caption)
        image_paragraph.add_run().add_picture(str(spec.path), width=Inches(spec.width_in))

    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(f".{destination.stem}.{uuid.uuid4().hex}.tmp.docx")
    try:
        document.save(temporary)
        _validate_saved_document(temporary, captions, figures, original_sections)
        os.replace(temporary, destination)
    finally:
        if temporary.exists():
            temporary.unlink()
    return destination


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--docx", type=Path, default=DOCX)
    parser.add_argument("--captions", type=Path, default=CAPTIONS_PATH)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args(argv)
    output = embed_figures(
        docx_path=args.docx,
        captions_path=args.captions,
        output_path=args.output,
    )
    print("saved", output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
