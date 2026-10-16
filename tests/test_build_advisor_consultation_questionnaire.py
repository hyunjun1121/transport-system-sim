"""TDD guards for the contextual, regenerable advisor-consultation DOCX."""

from __future__ import annotations

import json
import sys
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.build_advisor_consultation_questionnaire import (
    CONTENT_WIDTH_DXA,
    DEFAULT_CONTENT,
    KOREAN_FONT,
    build_questionnaire,
    load_content,
)


EXPECTED_ITEMS = [
    "수송 대상과 임무 완료 기준",
    "비교할 수송 방식과 자원 기준",
    "철도 이용 가능 범위",
    "버스와 운전자 확보 범위",
    "전시 교통과 장애 상황",
    "상황 파악과 대응 방식",
]

EXPECTED_CONTEXT_FRAGMENTS = (
    "전시 동원예비군 약 1,000명",
    "버스만 이용하는 방식과 철도·버스를 연계하는 방식",
    "열차를 1편만 쓸지 2편 이상 쓸지",
    "아래 항목을 군 관계자와 논의해 연구 가정으로 정하기 위한 자료다",
)

INTERNAL_ONLY_TERMS = (
    ".csv",
    ".json",
    ".yaml",
    ".py",
    "policy_id",
    "scenario_id",
    "mean_makespan",
    "CRN",
    "A1",
    "A2",
    "S→R",
    "600명",
    "30분 배차간격",
    "23대",
)

REMOVED_PACKAGING = (
    "자문 목적과 응답방법",
    "응답 기록 기준",
    "현재 모형과 자문 전 수정안",
    "1차 필수 자문항목",
    "2차 보충 자문항목",
    "토론용 기준·범위 초안",
    "최종 우선순위 질문",
    "회의 후 가정등록부",
    "자문 후 연구 반영 순서",
)


def _document_text(document) -> str:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_cells = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join(paragraphs + table_cells)


def _paragraph_numbering_format(document, paragraph) -> str:
    num_id = paragraph._p.pPr.numPr.numId.val
    numbering = document.part.numbering_part.element
    num = next(
        element
        for element in numbering.findall(qn("w:num"))
        if element.get(qn("w:numId")) == str(num_id)
    )
    abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    abstract = next(
        element
        for element in numbering.findall(qn("w:abstractNum"))
        if element.get(qn("w:abstractNumId")) == abstract_id
    )
    return abstract.find(qn("w:lvl")).find(qn("w:numFmt")).get(qn("w:val"))


def test_default_content_has_context_and_core_decisions_only() -> None:
    content = load_content(DEFAULT_CONTENT)
    assert content["schema_version"] == 3
    assert set(content) == {
        "schema_version",
        "editing_note",
        "title",
        "context",
        "items",
    }
    assert isinstance(content["context"], list)
    assert len(content["context"]) == 2
    assert all(isinstance(paragraph, str) and paragraph.strip() for paragraph in content["context"])
    context_text = "\n".join(content["context"])
    assert all(fragment in context_text for fragment in EXPECTED_CONTEXT_FRAGMENTS)
    assert [item["title"] for item in content["items"]] == EXPECTED_ITEMS
    assert len(content["items"]) == 6
    assert all(len(item["details"]) == 3 for item in content["items"])
    assert all(set(item) == {"id", "title", "details"} for item in content["items"])
    joined = "\n".join(
        [content["title"]]
        + [item["title"] for item in content["items"]]
        + [detail for item in content["items"] for detail in item["details"]]
    )
    assert not any(text in joined for text in REMOVED_PACKAGING)
    assert not any(text in joined for text in INTERNAL_ONLY_TERMS)


def test_builds_context_and_one_table_without_cover_or_page_furniture() -> None:
    with TemporaryDirectory() as directory:
        output = Path(directory) / "questionnaire.docx"
        built = build_questionnaire(DEFAULT_CONTENT, output)
        assert built == output.resolve()
        assert built.exists() and built.stat().st_size > 10_000

        document = Document(built)
        content = load_content(DEFAULT_CONTENT)
        text = _document_text(document)
        assert content["title"] in text
        assert all(paragraph in text for paragraph in content["context"])
        assert all(item in text for item in EXPECTED_ITEMS)
        assert not any(item in text for item in REMOVED_PACKAGING)
        assert len(document.tables) == 1
        table = document.tables[0]
        assert len(table.rows) == 7
        assert [cell.text for cell in table.rows[0].cells] == [
            "정해야 할 핵심 항목",
            "세부 결정내용",
        ]
        assert [row.cells[0].text for row in table.rows[1:]] == [
            f"{index}. {title}" for index, title in enumerate(EXPECTED_ITEMS, start=1)
        ]
        assert document.settings.element.find(qn("w:documentProtection")) is None
        assert not document.part.element.findall(".//" + qn("w:drawing"))
        assert not document.sections[0].header.paragraphs[0].text
        assert not document.sections[0].footer.paragraphs[0].text

        body_paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        assert body_paragraphs == [content["title"], *content["context"]]

        section = document.sections[0]
        assert int(round(section.page_width.twips)) == 12240
        assert int(round(section.page_height.twips)) == 15840
        assert int(round(section.left_margin.twips)) == 1440
        assert int(round(section.right_margin.twips)) == 1440

        normal = document.styles["Normal"]
        r_fonts = normal.element.rPr.find(qn("w:rFonts"))
        assert r_fonts is not None
        assert r_fonts.get(qn("w:eastAsia")) == KOREAN_FONT
        assert normal.font.size.pt == 11
        assert normal.paragraph_format.space_after.pt == 6

        grid = [int(column.get(qn("w:w"))) for column in table._tbl.tblGrid]
        assert grid == [2400, 6960]
        assert sum(grid) == CONTENT_WIDTH_DXA
        table_width = table._tbl.tblPr.find(qn("w:tblW"))
        assert table_width is not None
        assert table_width.get(qn("w:type")) == "dxa"
        assert int(table_width.get(qn("w:w"))) == CONTENT_WIDTH_DXA

        detail_paragraphs = [
            paragraph
            for row in table.rows[1:]
            for paragraph in row.cells[1].paragraphs
            if paragraph.text
        ]
        expected_detail_count = sum(len(item["details"]) for item in content["items"])
        assert len(detail_paragraphs) == expected_detail_count
        assert all(
            _paragraph_numbering_format(document, paragraph) == "bullet"
            for paragraph in detail_paragraphs
        )
        assert all(
            (paragraph.paragraph_format.space_after is None)
            or paragraph.paragraph_format.space_after.pt <= 1.0
            for paragraph in detail_paragraphs
        )

        for row in table.rows:
            for cell in row.cells:
                margins = cell._tc.tcPr.find(qn("w:tcMar"))
                assert margins is not None
                assert int(margins.find(qn("w:top")).get(qn("w:w"))) == 60
                assert int(margins.find(qn("w:bottom")).get(qn("w:w"))) == 60


def test_content_file_is_the_edit_surface() -> None:
    with TemporaryDirectory() as directory:
        root = Path(directory)
        content = json.loads(DEFAULT_CONTENT.read_text(encoding="utf-8"))
        replacement = "수정 가능한 연구 맥락 테스트"
        content["context"][0] = replacement
        content_path = root / "content.json"
        content_path.write_text(
            json.dumps(content, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        output = root / "modified.docx"
        build_questionnaire(content_path, output)
        assert replacement in _document_text(Document(output))


def test_overwrite_requires_explicit_force() -> None:
    with TemporaryDirectory() as directory:
        output = Path(directory) / "questionnaire.docx"
        build_questionnaire(DEFAULT_CONTENT, output)
        try:
            build_questionnaire(DEFAULT_CONTENT, output)
        except FileExistsError as exc:
            assert "--force" in str(exc)
        else:
            raise AssertionError("existing DOCX was overwritten without explicit force")
        build_questionnaire(DEFAULT_CONTENT, output, force=True)


def test_rejects_precise_coordinates() -> None:
    with TemporaryDirectory() as directory:
        content = json.loads(DEFAULT_CONTENT.read_text(encoding="utf-8"))
        content["items"][0]["details"][0] = "37.5202, 127.1210"
        path = Path(directory) / "unsafe.json"
        path.write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
        try:
            load_content(path)
        except ValueError as exc:
            assert "coordinate" in str(exc)
        else:
            raise AssertionError("precise coordinate pair was not rejected")


def test_rejects_internal_variable_or_file_references() -> None:
    with TemporaryDirectory() as directory:
        content = json.loads(DEFAULT_CONTENT.read_text(encoding="utf-8"))
        content["context"][0] = "세부 결과는 scenario_id와 results.csv를 참고한다."
        path = Path(directory) / "internal.json"
        path.write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
        try:
            load_content(path)
        except ValueError as exc:
            assert "internal-only" in str(exc)
        else:
            raise AssertionError("internal variable and file references were not rejected")


def main() -> None:
    test_default_content_has_context_and_core_decisions_only()
    test_builds_context_and_one_table_without_cover_or_page_furniture()
    test_content_file_is_the_edit_surface()
    test_overwrite_requires_explicit_force()
    test_rejects_precise_coordinates()
    test_rejects_internal_variable_or_file_references()
    print("minimal advisor consultation questionnaire tests: pass")


if __name__ == "__main__":
    main()
