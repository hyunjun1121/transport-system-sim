from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from paper._build.update_revision_manuscript import (
    AnalysisIntegrityError,
    AnchorNotFoundError,
    AnchorNotUniqueError,
    ProtectedContentError,
    atomic_save_document,
    build_document_inventory,
    find_table_after_anchor,
    find_unique_paragraph,
    load_verified_analysis,
    open_manuscript_session,
    replace_paragraph_text,
    replace_table_rows,
)
from paper._build.manuscript_content_plan import (
    ContentPlanError,
    apply_content_plan,
    load_content_plan,
    preview_content_plan,
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def _analysis_fixture(root: Path) -> Path:
    output_root = root / "paper_revision_top5_20260721"
    analysis_dir = output_root / "analysis" / "full"
    analysis_dir.mkdir(parents=True)
    paired_path = analysis_dir / "paired_summary.csv"
    _write_csv(
        paired_path,
        [
            {
                "scenario_id": "no_disruption",
                "total_pair_count": 30,
                "mean_bus_makespan": 283.1,
                "mean_multimodal_makespan": 363.8,
                "mean_delta_bus_minus_multimodal": -80.7,
                "completion_t_ci_lower": -0.001,
                "completion_t_ci_upper": 0.001,
                "incomplete_reason_counts": "{}",
            }
        ],
    )
    adaptive_path = analysis_dir / "adaptive_policies.csv"
    _write_csv(
        adaptive_path,
        [
            {
                "policy_id": "precheck_switch",
                "run_count": 30,
                "mean_completion_rate": 1.0,
                "winner": True,
            }
        ],
    )
    source_path = output_root / "paired_reanalysis" / "full_results.csv"
    _write_csv(source_path, [{"campaign_id": "paired_reanalysis", "run_key": "a" * 64}])
    generated = []
    for path in (paired_path, adaptive_path):
        generated.append(
            {
                "path": path.name,
                "sha256": _sha256(path),
                "data_row_count": 1,
            }
        )
    manifest = {
        "schema_version": 2,
        "final_study_ready": False,
        "stage": "full",
        "analysis_complete_full": True,
        "analysis_completeness_blockers": [],
        "output_root": str(output_root.resolve()),
        "analysis_directory": str(analysis_dir.resolve()),
        "source_files": [
            {
                "path": "paired_reanalysis/full_results.csv",
                "sha256": _sha256(source_path),
                "data_row_count": 1,
            }
        ],
        "analyses": {
            "paired_summary": {"status": "available"},
            "adaptive_policies": {"status": "available"},
        },
        "generated_csv_files": generated,
    }
    manifest_path = analysis_dir / "analysis_manifest.json"
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return manifest_path


def _document_fixture(path: Path, image_path: Path | None = None) -> None:
    document = Document()
    for text in ("제목", "부제", "영문 제목", "저자 A", "저자 B", "저자 C"):
        document.add_paragraph(text)
    heading = document.add_paragraph("1. 서론", style="Heading 1")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(7)
    run = heading.runs[0]
    run.font.name = "Arial"
    run.font.size = Pt(13)
    document.add_paragraph("본문 앵커")
    document.add_paragraph("<표 1> 결과")
    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.columns[0].width = Inches(1.25)
    table.columns[1].width = Inches(2.5)
    table.cell(0, 0).text = "항목"
    table.cell(0, 1).text = "값"
    table.cell(1, 0).text = "기존"
    table.cell(1, 1).text = "1"
    if image_path is not None:
        paragraph = document.add_paragraph()
        paragraph.add_run().add_picture(str(image_path), width=Inches(0.2))
    document.save(path)


def _tiny_png(path: Path) -> None:
    path.write_bytes(
        bytes.fromhex(
            "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
            "0000000d49444154789c6360f8cff00000040101005f5dc60b0000000049454e44ae426082"
        )
    )


def test_analysis_manifest_and_typed_table_guards() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        manifest_path = _analysis_fixture(Path(temp_dir))
        bundle = load_verified_analysis(manifest_path)
        paired = bundle.table("paired_summary")
        row = paired.one(scenario_id="no_disruption")
        assert row["total_pair_count"] == 30
        assert isinstance(row["total_pair_count"], int)
        assert row["mean_delta_bus_minus_multimodal"] == -80.7
        assert isinstance(row["mean_delta_bus_minus_multimodal"], float)
        assert row["mean_bus_makespan"] == 283.1
        assert isinstance(row["mean_bus_makespan"], float)
        assert row["mean_multimodal_makespan"] == 363.8
        assert isinstance(row["mean_multimodal_makespan"], float)
        assert row["incomplete_reason_counts"] == {}
        adaptive = bundle.table("adaptive_policies.csv")
        assert adaptive.rows[0]["winner"] is True

        paired.path.write_text("tampered\n", encoding="utf-8")
        try:
            load_verified_analysis(manifest_path)
        except AnalysisIntegrityError as error:
            assert "SHA-256" in str(error)
        else:
            raise AssertionError("tampered analysis CSV was accepted")

    with tempfile.TemporaryDirectory() as temp_dir:
        manifest_path = _analysis_fixture(Path(temp_dir))
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["stage"] = "smoke"
        manifest["analysis_complete_full"] = False
        manifest["analysis_completeness_blockers"] = [
            "analysis_stage_is_not_full"
        ]
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        try:
            load_verified_analysis(manifest_path)
        except AnalysisIntegrityError as error:
            assert "stage=full" in str(error)
        else:
            raise AssertionError("smoke analysis was accepted")
    print("PASS: full-analysis integrity and typed CSV readers")


def test_unique_anchor_and_paragraph_format_preservation() -> None:
    document = Document()
    paragraph = document.add_paragraph("Purpose. Old", style="Heading 1")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.runs[0].font.name = "Arial"
    paragraph.runs[0].font.size = Pt(12)
    ppr_before = paragraph._p.pPr.xml
    style_before = paragraph.style.style_id

    found = find_unique_paragraph(document, "Purpose. Old")
    assert found._p is paragraph._p
    replace_paragraph_text(
        found,
        "Purpose. New evidence statement.",
        first_run_label="Purpose.",
        first_run_bold=True,
    )
    assert found.text == "Purpose. New evidence statement."
    assert found.style.style_id == style_before
    assert found._p.pPr.xml == ppr_before
    assert len(found.runs) == 2
    assert found.runs[0].text == "Purpose."
    assert found.runs[0].bold is True
    assert found.runs[0].font.name == "Arial"
    assert found.runs[0].font.size == Pt(12)

    document.add_paragraph("중복")
    document.add_paragraph("중복")
    try:
        find_unique_paragraph(document, "중복")
    except AnchorNotUniqueError:
        pass
    else:
        raise AssertionError("duplicate paragraph anchor was accepted")
    try:
        find_unique_paragraph(document, "없음")
    except AnchorNotFoundError:
        pass
    else:
        raise AssertionError("missing paragraph anchor was accepted")
    print("PASS: unique anchors and paragraph formatting preservation")


def test_table_lookup_and_geometry_preservation() -> None:
    document = Document()
    document.add_paragraph("<표 1> 결과")
    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "항목"
    table.cell(0, 1).text = "값"
    table.cell(1, 0).text = "기존"
    table.cell(1, 1).text = "1"
    for cell, width in zip(table.rows[0].cells, (1800, 3600)):
        tc_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
        if tc_width is None:
            tc_width = OxmlElement("w:tcW")
            cell._tc.get_or_add_tcPr().append(tc_width)
        tc_width.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w", str(width))
        tc_width.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type", "dxa")
    located = find_table_after_anchor(document, "<표 1> 결과")
    assert located._tbl is table._tbl
    tbl_pr_before = table._tbl.tblPr.xml
    tbl_grid_before = table._tbl.tblGrid.xml
    header_before = table.rows[0]._tr.xml
    tc_pr_before = [cell._tc.tcPr.xml for cell in table.rows[1].cells]

    replace_table_rows(
        table,
        [("A", 1.25), ("B", 2.5), ("C", 3.75)],
        preserve_leading_rows=1,
    )
    assert len(table.rows) == 4
    assert [cell.text for cell in table.rows[0].cells] == ["항목", "값"]
    assert [cell.text for cell in table.rows[3].cells] == ["C", "3.75"]
    assert table._tbl.tblPr.xml == tbl_pr_before
    assert table._tbl.tblGrid.xml == tbl_grid_before
    assert table.rows[0]._tr.xml == header_before
    assert [cell._tc.tcPr.xml for cell in table.rows[1].cells] == tc_pr_before

    replace_table_rows(table, [("D", 4)], preserve_leading_rows=1)
    assert len(table.rows) == 2
    assert [cell.text for cell in table.rows[1].cells] == ["D", "4"]
    assert table._tbl.tblPr.xml == tbl_pr_before
    assert table._tbl.tblGrid.xml == tbl_grid_before
    print("PASS: table lookup, row resize, and XML geometry preservation")


def test_atomic_save_preserves_protected_content_drawings_and_sections() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        image_path = root / "pixel.png"
        _tiny_png(image_path)
        source = root / "source.docx"
        output = root / "output.docx"
        _document_fixture(source, image_path)

        session = open_manuscript_session(source)
        replace_paragraph_text(
            find_unique_paragraph(session.document, "본문 앵커"), "수정 본문"
        )
        reopened = atomic_save_document(session, output)
        assert output.exists()
        assert [reopened.paragraphs[index].text for index in (3, 4, 5)] == [
            "저자 A",
            "저자 B",
            "저자 C",
        ]
        assert len(reopened.inline_shapes) == 1
        assert len(reopened.sections) == 1
        assert reopened.styles.element.xml == Document(source).styles.element.xml

        failed_session = open_manuscript_session(source)
        failed_session.document.paragraphs[3].text = "변조"
        original_output = output.read_bytes()
        try:
            atomic_save_document(failed_session, output)
        except ProtectedContentError as error:
            assert "protected author" in str(error)
        else:
            raise AssertionError("protected author mutation was saved")
        assert output.read_bytes() == original_output
        assert not list(root.glob(".output.*.tmp.docx"))

        style_session = open_manuscript_session(source)
        style_session.document.styles["Normal"].font.name = "Comic Sans MS"
        original_output = output.read_bytes()
        try:
            atomic_save_document(style_session, output)
        except ProtectedContentError as error:
            assert "style definitions" in str(error)
        else:
            raise AssertionError("global style mutation was saved")
        assert output.read_bytes() == original_output
        assert not list(root.glob(".output.*.tmp.docx"))
    print("PASS: atomic save and protected-content invariants")


def test_dry_run_inventory_is_read_only() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        source = root / "source.docx"
        _document_fixture(source)
        before = source.read_bytes()
        session = open_manuscript_session(source)
        inventory = build_document_inventory(
            session.document, anchors=("1. 서론", "<표 1> 결과")
        )
        assert inventory["paragraph_count"] == 9
        assert inventory["table_count"] == 1
        assert inventory["section_count"] == 1
        assert inventory["anchors"]["1. 서론"]["match_count"] == 1
        assert inventory["anchors"]["1. 서론"]["following_table"] is False
        assert inventory["anchors"]["<표 1> 결과"]["following_table"] is True
        assert source.read_bytes() == before
    print("PASS: dry-run inventory is read-only")


def _write_content_plan(
    path: Path,
    *,
    source_path: Path,
    analysis_manifest_path: Path,
    operations: list[dict[str, object]] | None = None,
) -> None:
    if operations is None:
        operations = [
            {
                "id": "body-result",
                "kind": "paragraph",
                "anchor": {"text": "본문 앵커", "mode": "exact"},
                "replacement": {
                    "fragments": [
                        {"literal": "분석값 "},
                        {
                            "evidence": {
                                "table": "paired_summary",
                                "where": {"scenario_id": "no_disruption"},
                                "column": "mean_delta_bus_minus_multimodal",
                                "format": ".1f",
                            }
                        },
                    ]
                },
            },
            {
                "id": "result-table",
                "kind": "table_rows",
                "anchor": {"text": "<표 1> 결과", "mode": "exact"},
                "preserve_leading_rows": 1,
                "header": ["정책", "평균 완료율"],
                "source": {
                    "table": "adaptive_policies",
                    "where": {"winner": True},
                    "sort_by": ["policy_id"],
                },
                "cells": [
                    {"column": "policy_id"},
                    {"column": "mean_completion_rate", "format": ".3f"},
                ],
            },
        ]
    payload = {
        "schema_version": 1,
        "plan_id": "fixture-revision",
        "source_sha256": _sha256(source_path),
        "analysis_manifest_sha256": _sha256(analysis_manifest_path),
        "operations": operations,
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def test_declarative_plan_preview_and_atomic_apply() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        image_path = root / "pixel.png"
        _tiny_png(image_path)
        source = root / "source.docx"
        output = root / "output.docx"
        plan_path = root / "content_plan.json"
        _document_fixture(source, image_path)
        analysis_manifest = _analysis_fixture(root)
        _write_content_plan(
            plan_path,
            source_path=source,
            analysis_manifest_path=analysis_manifest,
        )
        source_before = source.read_bytes()

        plan = load_content_plan(plan_path)
        assert plan.plan_id == "fixture-revision"
        preview = preview_content_plan(source, analysis_manifest, plan_path)
        assert preview["dry_run"] is True
        assert preview["operation_count"] == 2
        assert preview["operations"][0]["resolved_text"] == "분석값 -80.7"
        assert preview["operations"][1]["resolved_row_count"] == 1
        assert preview["operations"][1]["resolved_rows_preview"] == [
            ["precheck_switch", "1.000"]
        ]
        assert preview["operations"][1]["header"] == ["정책", "평균 완료율"]
        assert source.read_bytes() == source_before
        assert not output.exists()

        report = apply_content_plan(source, analysis_manifest, plan_path, output)
        assert report["dry_run"] is False
        assert report["operation_count"] == 2
        assert report["output_path"] == str(output.resolve())
        assert report["output_sha256"] == _sha256(output)
        assert source.read_bytes() == source_before
        revised = Document(output)
        assert revised.paragraphs[7].text == "분석값 -80.7"
        assert [revised.paragraphs[index].text for index in (3, 4, 5)] == [
            "저자 A",
            "저자 B",
            "저자 C",
        ]
        assert len(revised.inline_shapes) == 1
        assert len(revised.sections) == 1
        assert [[cell.text for cell in row.cells] for row in revised.tables[0].rows] == [
            ["정책", "평균 완료율"],
            ["precheck_switch", "1.000"],
        ]
    print("PASS: declarative plan preview and atomic apply")


def test_content_plan_applies_pagination_controls_before_atomic_save() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        source = root / "source.docx"
        output = root / "output.docx"
        plan_path = root / "content_plan.json"
        _document_fixture(source)
        source_document = Document(source)
        source_document.add_paragraph("8.3 추가 분석")
        source_document.save(source)
        analysis_manifest = _analysis_fixture(root)
        _write_content_plan(
            plan_path,
            source_path=source,
            analysis_manifest_path=analysis_manifest,
        )

        apply_content_plan(source, analysis_manifest, plan_path, output)
        revised = Document(output)
        table = revised.tables[0]

        assert all(
            len(row._tr.xpath("./w:trPr/w:cantSplit")) == 1
            for row in table.rows
        )
        assert len(table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")) == 1
        assert all(
            not row._tr.xpath("./w:trPr/w:tblHeader")
            for row in table.rows[1:]
        )
        assert all(
            paragraph.paragraph_format.keep_with_next is True
            for cell in table.rows[0].cells
            for paragraph in cell.paragraphs
        )

        caption = next(
            paragraph
            for paragraph in revised.paragraphs
            if paragraph.text.startswith("<표 ")
        )
        heading = next(
            paragraph
            for paragraph in revised.paragraphs
            if paragraph.text.startswith("8.3 ")
        )
        body = next(
            paragraph
            for paragraph in revised.paragraphs
            if paragraph.text == "분석값 -80.7"
        )
        assert caption.paragraph_format.keep_with_next is True
        assert heading.paragraph_format.keep_with_next is True
        assert body.paragraph_format.keep_with_next is not True
    print("PASS: production content plan applies pagination controls")


def test_plan_hash_guards_and_preflight_are_transactional() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        source = root / "source.docx"
        output = root / "existing.docx"
        plan_path = root / "content_plan.json"
        _document_fixture(source)
        analysis_manifest = _analysis_fixture(root)
        output.write_bytes(b"existing-output")
        original_output = output.read_bytes()

        _write_content_plan(
            plan_path,
            source_path=source,
            analysis_manifest_path=analysis_manifest,
        )
        payload = json.loads(plan_path.read_text(encoding="utf-8"))
        payload["source_sha256"] = "0" * 64
        plan_path.write_text(json.dumps(payload), encoding="utf-8")
        try:
            preview_content_plan(source, analysis_manifest, plan_path)
        except ContentPlanError as error:
            assert "source_sha256" in str(error)
        else:
            raise AssertionError("wrong source hash was accepted")
        assert output.read_bytes() == original_output

        invalid_operations = [
            {
                "id": "valid-first",
                "kind": "paragraph",
                "anchor": {"text": "본문 앵커"},
                "replacement": {"fragments": [{"literal": "수정 시도"}]},
            },
            {
                "id": "invalid-second",
                "kind": "paragraph",
                "anchor": {"text": "없는 앵커"},
                "replacement": {"fragments": [{"literal": "도달 금지"}]},
            },
        ]
        _write_content_plan(
            plan_path,
            source_path=source,
            analysis_manifest_path=analysis_manifest,
            operations=invalid_operations,
        )
        source_before = source.read_bytes()
        try:
            apply_content_plan(source, analysis_manifest, plan_path, output)
        except AnchorNotFoundError:
            pass
        else:
            raise AssertionError("plan with missing second anchor was applied")
        assert source.read_bytes() == source_before
        assert output.read_bytes() == original_output
    print("PASS: plan hashes and transactional preflight")


def test_declarative_plan_accepts_literal_table_rows() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        source = root / "source.docx"
        output = root / "output.docx"
        plan_path = root / "content_plan.json"
        _document_fixture(source)
        analysis_manifest = _analysis_fixture(root)
        operations = [
            {
                "id": "literal-result-table",
                "kind": "table_rows",
                "anchor": {"text": "<표 1> 결과"},
                "preserve_leading_rows": 1,
                "header": ["구분", "결과"],
                "rows": [["기준선", "조건부"], ["경계", "1.742"]],
            }
        ]
        _write_content_plan(
            plan_path,
            source_path=source,
            analysis_manifest_path=analysis_manifest,
            operations=operations,
        )

        preview = preview_content_plan(source, analysis_manifest, plan_path)
        assert preview["operations"][0]["resolved_rows_preview"] == [
            ["기준선", "조건부"],
            ["경계", "1.742"],
        ]
        apply_content_plan(source, analysis_manifest, plan_path, output)
        revised = Document(output)
        assert [[cell.text for cell in row.cells] for row in revised.tables[0].rows] == [
            ["구분", "결과"],
            ["기준선", "조건부"],
            ["경계", "1.742"],
        ]
    print("PASS: declarative plan accepts literal table rows")


def test_declarative_plan_sorts_numeric_columns_numerically() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        source = root / "source.docx"
        plan_path = root / "content_plan.json"
        _document_fixture(source)
        analysis_manifest = _analysis_fixture(root)
        analysis_dir = analysis_manifest.parent
        adaptive_path = analysis_dir / "adaptive_policies.csv"
        _write_csv(
            adaptive_path,
            [
                {
                    "policy_id": "rank_two",
                    "run_count": 30,
                    "mean_completion_rate": 1.0,
                    "winner": False,
                    "rank": 2,
                },
                {
                    "policy_id": "rank_ten",
                    "run_count": 30,
                    "mean_completion_rate": 1.0,
                    "winner": False,
                    "rank": 10,
                },
                {
                    "policy_id": "rank_one",
                    "run_count": 30,
                    "mean_completion_rate": 1.0,
                    "winner": True,
                    "rank": 1,
                },
            ],
        )
        manifest = json.loads(analysis_manifest.read_text(encoding="utf-8"))
        adaptive_record = next(
            item
            for item in manifest["generated_csv_files"]
            if item["path"] == "adaptive_policies.csv"
        )
        adaptive_record["sha256"] = _sha256(adaptive_path)
        adaptive_record["data_row_count"] = 3
        analysis_manifest.write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        operations = [
            {
                "id": "numeric-sort-table",
                "kind": "table_rows",
                "anchor": {"text": "<표 1> 결과"},
                "preserve_leading_rows": 1,
                "header": ["정책", "순위"],
                "source": {
                    "table": "adaptive_policies",
                    "where": {},
                    "sort_by": ["rank"],
                },
                "cells": [
                    {"column": "policy_id"},
                    {"column": "rank"},
                ],
            }
        ]
        _write_content_plan(
            plan_path,
            source_path=source,
            analysis_manifest_path=analysis_manifest,
            operations=operations,
        )

        preview = preview_content_plan(source, analysis_manifest, plan_path)

        assert preview["operations"][0]["resolved_rows_preview"] == [
            ["rank_one", "1"],
            ["rank_two", "2"],
            ["rank_ten", "10"],
        ]
    print("PASS: declarative plan sorts numeric columns numerically")


def test_plan_rejects_ambiguous_or_unsupported_evidence() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        source = root / "source.docx"
        plan_path = root / "content_plan.json"
        _document_fixture(source)
        analysis_manifest = _analysis_fixture(root)

        duplicate_target = {
            "id": "duplicate-target",
            "kind": "paragraph",
            "anchor": {"text": "본문 앵커"},
            "replacement": {"fragments": [{"literal": "두 번째"}]},
        }
        operations = [
            {
                "id": "first-target",
                "kind": "paragraph",
                "anchor": {"text": "본문 앵커"},
                "replacement": {"fragments": [{"literal": "첫 번째"}]},
            },
            duplicate_target,
        ]
        _write_content_plan(
            plan_path,
            source_path=source,
            analysis_manifest_path=analysis_manifest,
            operations=operations,
        )
        try:
            preview_content_plan(source, analysis_manifest, plan_path)
        except ContentPlanError as error:
            assert "duplicate target" in str(error)
        else:
            raise AssertionError("duplicate plan target was accepted")

        operations = [
            {
                "id": "bad-evidence",
                "kind": "paragraph",
                "anchor": {"text": "본문 앵커"},
                "replacement": {
                    "fragments": [
                        {
                            "evidence": {
                                "table": "paired_summary",
                                "where": {},
                                "column": "missing_column",
                            }
                        }
                    ]
                },
            }
        ]
        _write_content_plan(
            plan_path,
            source_path=source,
            analysis_manifest_path=analysis_manifest,
            operations=operations,
        )
        try:
            preview_content_plan(source, analysis_manifest, plan_path)
        except ContentPlanError as error:
            assert "missing_column" in str(error)
        else:
            raise AssertionError("unknown evidence column was accepted")
    print("PASS: ambiguous targets and unsupported evidence are rejected")


def test_content_plan_cli_is_directly_executable() -> None:
    script = ROOT / "paper" / "_build" / "manuscript_content_plan.py"
    completed = subprocess.run(
        [sys.executable, str(script), "--help"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr
    assert "--analysis-manifest" in completed.stdout
    assert "--dry-run" in completed.stdout
    print("PASS: content-plan CLI is directly executable")


def test_production_plan_binds_core_claims_and_excludes_audited_mismatches() -> None:
    path = ROOT / "paper" / "_build" / "revision_content_plan_v4.json"
    payload = json.loads(path.read_text(encoding="utf-8"))
    operations = {item["id"]: item for item in payload["operations"]}

    def evidence_count(operation_id: str) -> int:
        fragments = operations[operation_id]["replacement"]["fragments"]
        return sum("evidence" in fragment for fragment in fragments)

    assert evidence_count("results-baseline") >= 14
    assert evidence_count("results-boundary") >= 3
    serialized = json.dumps(payload, ensure_ascii=False)
    for stale in (
        "도로 10수준×철도 3상태",
        "top-3는 정상 최단시간을 보존",
        "342.5",
        "442·472·502",
        "시간대별 누적도착",
        "STRICT와 GRACE는 도착",
    ):
        assert stale not in serialized
    print("PASS: production plan binds core claims and excludes audited mismatches")


if __name__ == "__main__":
    test_analysis_manifest_and_typed_table_guards()
    test_unique_anchor_and_paragraph_format_preservation()
    test_table_lookup_and_geometry_preservation()
    test_atomic_save_preserves_protected_content_drawings_and_sections()
    test_dry_run_inventory_is_read_only()
    test_declarative_plan_preview_and_atomic_apply()
    test_content_plan_applies_pagination_controls_before_atomic_save()
    test_plan_hash_guards_and_preflight_are_transactional()
    test_declarative_plan_accepts_literal_table_rows()
    test_declarative_plan_sorts_numeric_columns_numerically()
    test_plan_rejects_ambiguous_or_unsupported_evidence()
    test_content_plan_cli_is_directly_executable()
    test_production_plan_binds_core_claims_and_excludes_audited_mismatches()
    print("\n=== REVISION MANUSCRIPT UPDATE TESTS PASSED ===")
