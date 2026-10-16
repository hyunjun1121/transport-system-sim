"""TDD guards for targeted, caption-driven DOCX figure embedding."""

from __future__ import annotations

import hashlib
import importlib.util
import sys
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
MODULE_PATH = ROOT / "paper" / "figures" / "_build" / "embed_figures.py"


def _load_module():
    spec = importlib.util.spec_from_file_location("embed_figures_for_test", MODULE_PATH)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _write_png(path: Path, rgb: tuple[int, int, int]) -> None:
    Image.new("RGB", (24, 16), rgb).save(path)


def _has_drawing(paragraph) -> bool:
    return bool(paragraph._p.findall(".//" + qn("w:drawing")))


def _drawing_blob_sha(paragraph) -> str:
    blips = paragraph._p.xpath(".//a:blip")
    assert len(blips) == 1
    rel_id = blips[0].get(qn("r:embed"))
    assert rel_id
    return hashlib.sha256(paragraph.part.related_parts[rel_id].blob).hexdigest()


def _paragraph_before(paragraph):
    previous = paragraph._p.getprevious()
    assert previous is not None and previous.tag == qn("w:p")
    return next(p for p in paragraph._parent.paragraphs if p._p is previous)


def _write_captions(path: Path) -> dict[str, str]:
    lines = ["# 테스트 캡션", "", "설명 문단은 캡션이 아니다.", "", "---", ""]
    expected = {}
    for number in range(1, 7):
        body = f"캡션 원본 {number}: 첫 문장. 둘째 문장 {number}."
        lines.extend([f"**그림 {number}.** 캡션 원본 {number}: 첫 문장.", f"둘째 문장 {number}.", ""])
        expected[f"<그림 {number}>"] = f"<그림 {number}> {body}"
    path.write_text("\n".join(lines), encoding="utf-8")
    return expected


def _make_fixture(root: Path, module):
    captions_path = root / "captions.md"
    expected = _write_captions(captions_path)

    unrelated = root / "unrelated.png"
    old_target = root / "old-target.png"
    _write_png(unrelated, (5, 10, 15))
    _write_png(old_target, (40, 45, 50))

    specs = []
    for number in range(1, 7):
        image_path = root / f"figure-{number}.png"
        _write_png(image_path, (20 * number, 10 * number, 5 * number))
        specs.append(module.FigureSpec(number, image_path, 1.4))

    docx_path = root / "fixture.docx"
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Inches(0.73)

    unrelated_paragraph = document.add_paragraph()
    unrelated_paragraph.add_run("UNRELATED-DRAWING-GUARD ")
    unrelated_paragraph.add_run().add_picture(str(unrelated), width=Inches(0.3))
    document.add_paragraph("ordinary text keeps this drawing unrelated to every caption")

    for number in range(1, 7):
        image_paragraph = document.add_paragraph(style="Normal")
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        image_paragraph.paragraph_format.left_indent = Pt(11 + number)
        image_paragraph.paragraph_format.space_after = Pt(2 + number)
        image_paragraph.add_run().add_picture(str(old_target), width=Inches(0.25))

        caption = document.add_paragraph(style="Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        caption.paragraph_format.left_indent = Pt(20 + number)
        run = caption.add_run(f"<그림 {number}> 이전 캡션")
        run.bold = True

    document.save(docx_path)
    return docx_path, captions_path, expected, specs, unrelated


def test_targeted_embedding_preserves_unrelated_drawing_and_formatting() -> None:
    module = _load_module()
    with TemporaryDirectory() as tmp:
        root = Path(tmp)
        docx_path, captions_path, expected, specs, unrelated = _make_fixture(root, module)
        original = Document(docx_path)
        original_section = (
            len(original.sections),
            original.sections[0].orientation,
            original.sections[0].page_width,
            original.sections[0].page_height,
            original.sections[0].left_margin,
        )
        target_format = {}
        caption_format = {}
        for number in range(1, 7):
            caption = next(p for p in original.paragraphs if p.text.startswith(f"<그림 {number}>"))
            image_paragraph = _paragraph_before(caption)
            target_format[number] = (
                image_paragraph.style.name,
                image_paragraph.alignment,
                image_paragraph.paragraph_format.left_indent,
                image_paragraph.paragraph_format.space_after,
            )
            caption_format[number] = (
                caption.style.name,
                caption.alignment,
                caption.paragraph_format.left_indent,
            )

        module.embed_figures(
            docx_path=docx_path,
            captions_path=captions_path,
            output_path=docx_path,
            figures=specs,
        )

        result = Document(docx_path)
        unrelated_paragraph = next(p for p in result.paragraphs if "UNRELATED-DRAWING-GUARD" in p.text)
        assert _has_drawing(unrelated_paragraph)
        assert _drawing_blob_sha(unrelated_paragraph) == hashlib.sha256(unrelated.read_bytes()).hexdigest()

        for number, spec in enumerate(specs, start=1):
            captions = [p for p in result.paragraphs if p.text.startswith(f"<그림 {number}>")]
            assert len(captions) == 1
            caption = captions[0]
            assert caption.text == expected[f"<그림 {number}>"]
            image_paragraph = _paragraph_before(caption)
            assert _has_drawing(image_paragraph)
            assert _drawing_blob_sha(image_paragraph) == hashlib.sha256(Path(spec.path).read_bytes()).hexdigest()
            assert (
                image_paragraph.style.name,
                image_paragraph.alignment,
                image_paragraph.paragraph_format.left_indent,
                image_paragraph.paragraph_format.space_after,
            ) == target_format[number]
            assert (
                caption.style.name,
                caption.alignment,
                caption.paragraph_format.left_indent,
            ) == caption_format[number]

        assert (
            len(result.sections),
            result.sections[0].orientation,
            result.sections[0].page_width,
            result.sections[0].page_height,
            result.sections[0].left_margin,
        ) == original_section
        assert not list(root.glob(".*.tmp.docx"))


def test_mixed_text_and_drawing_before_caption_is_preserved() -> None:
    module = _load_module()
    with TemporaryDirectory() as tmp:
        root = Path(tmp)
        docx_path, captions_path, expected, specs, _unrelated = _make_fixture(root, module)
        document = Document(docx_path)
        caption = next(p for p in document.paragraphs if p.text.startswith("<그림 1>"))
        mixed = _paragraph_before(caption)
        mixed.add_run("KEEP-TEXT")
        old_drawing_sha = _drawing_blob_sha(mixed)
        document.save(docx_path)

        module.embed_figures(
            docx_path=docx_path,
            captions_path=captions_path,
            output_path=docx_path,
            figures=specs,
        )

        result = Document(docx_path)
        caption = next(p for p in result.paragraphs if p.text.startswith("<그림 1>"))
        image_paragraph = _paragraph_before(caption)
        preserved = _paragraph_before(image_paragraph)
        assert caption.text == expected["<그림 1>"]
        assert image_paragraph.text == ""
        assert _drawing_blob_sha(image_paragraph) == hashlib.sha256(specs[0].path.read_bytes()).hexdigest()
        assert preserved.text == "KEEP-TEXT"
        assert _drawing_blob_sha(preserved) == old_drawing_sha


def test_duplicate_figure_numbers_are_rejected_before_write() -> None:
    module = _load_module()
    with TemporaryDirectory() as tmp:
        root = Path(tmp)
        docx_path, captions_path, _expected, specs, _unrelated = _make_fixture(root, module)
        before = docx_path.read_bytes()
        duplicate_specs = [specs[0], module.FigureSpec(1, specs[1].path, 1.4)]
        try:
            module.embed_figures(
                docx_path=docx_path,
                captions_path=captions_path,
                output_path=docx_path,
                figures=duplicate_specs,
            )
        except ValueError as exc:
            assert str(exc) == "duplicate figure number: 1"
        else:
            raise AssertionError("duplicate figure number was accepted")
        assert docx_path.read_bytes() == before
        assert not list(root.glob(".*.tmp.docx"))


def test_saved_document_validation_checks_source_image_blob() -> None:
    module = _load_module()
    with TemporaryDirectory() as tmp:
        root = Path(tmp)
        docx_path, captions_path, expected, specs, _unrelated = _make_fixture(root, module)
        module.embed_figures(
            docx_path=docx_path,
            captions_path=captions_path,
            output_path=docx_path,
            figures=specs,
        )
        saved = Document(docx_path)
        signature = module._section_signature(saved)
        _write_png(specs[0].path, (231, 17, 93))
        try:
            module._validate_saved_document(docx_path, expected, specs, signature)
        except RuntimeError as exc:
            assert str(exc) == "image blob mismatch after save: <그림 1>"
        else:
            raise AssertionError("source image blob mismatch was accepted")


def test_atomic_validation_failure_leaves_original_unchanged() -> None:
    module = _load_module()
    with TemporaryDirectory() as tmp:
        root = Path(tmp)
        docx_path, captions_path, _expected, specs, _unrelated = _make_fixture(root, module)
        before = docx_path.read_bytes()
        original_validator = module._validate_saved_document
        module._validate_saved_document = lambda *_args, **_kwargs: (_ for _ in ()).throw(
            RuntimeError("forced validation failure")
        )
        try:
            try:
                module.embed_figures(
                    docx_path=docx_path,
                    captions_path=captions_path,
                    output_path=docx_path,
                    figures=specs,
                )
            except RuntimeError as exc:
                assert str(exc) == "forced validation failure"
            else:
                raise AssertionError("validation failure did not propagate")
        finally:
            module._validate_saved_document = original_validator
        assert docx_path.read_bytes() == before
        assert not list(root.glob(".*.tmp.docx"))


def main() -> int:
    test_targeted_embedding_preserves_unrelated_drawing_and_formatting()
    print("PASS: targeted embedding preserves unrelated drawing and formatting")
    test_mixed_text_and_drawing_before_caption_is_preserved()
    print("PASS: mixed text/drawing paragraph is preserved")
    test_duplicate_figure_numbers_are_rejected_before_write()
    print("PASS: duplicate figure numbers are rejected before write")
    test_saved_document_validation_checks_source_image_blob()
    print("PASS: saved drawing blob matches source image")
    test_atomic_validation_failure_leaves_original_unchanged()
    print("PASS: failed validation leaves original DOCX unchanged")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
